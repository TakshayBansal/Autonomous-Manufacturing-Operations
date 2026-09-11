"""Safe Excel/CSV external-system connector.

Parsing never evaluates workbook content. Preview persists source lineage and is
the only input accepted by commit, keeping manual and assistant imports on the
same governed path.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import secrets
import zipfile
from datetime import datetime, timezone
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.db import models
from app.domains import workflows


MAX_SHEETS = 40
MAX_ROWS = 100_000
MAX_ARCHIVE_FILES = 50
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
FORMULA_PREFIXES = ("=", "+", "-", "@")

SHEET_SCHEMAS: dict[str, dict[str, Any]] = {
    "Organizations": {"entity": "company", "key": "external_key", "required": ["external_key", "name"]},
    "Plants": {"entity": "plant", "key": "external_key", "required": ["external_key", "organization_key", "code", "name"]},
    "Users": {"entity": "user", "key": "external_key", "required": ["external_key", "name", "email", "role", "plant_key"]},
    "ReportingLines": {"entity": "reporting_line", "key": "employee_key", "required": ["employee_key", "manager_key"]},
    "Items": {"entity": "item", "key": "external_key", "required": ["external_key", "code", "name", "uom"]},
    "Suppliers": {"entity": "supplier", "key": "external_key", "required": ["external_key", "name"]},
    "ItemSpecifications": {"entity": "item_specification", "key": "external_key", "required": ["external_key", "item_external_key", "description"]},
    "SupplierContacts": {"entity": "supplier_contact", "key": "external_key", "required": ["external_key", "supplier_external_key", "name", "email"]},
    "SupplierCapabilities": {"entity": "supplier_item_capability", "key": "external_key", "required": ["external_key", "supplier_external_key", "item_external_key", "approved"]},
    "Requirements": {"entity": "purchase_requirement", "key": "external_key", "required": ["external_key", "item_code", "quantity", "uom", "need_by_date", "reason"]},
    "OpenPOs": {"entity": "po_draft", "key": "external_key", "required": ["external_key", "supplier_external_key", "status"]},
    "OpenPOLines": {"entity": "po_draft_line", "key": "external_key", "required": ["external_key", "po_external_key", "item_external_key", "quantity", "uom", "unit_price", "need_by_date"]},
    "Receipts": {"entity": "store_receipt", "key": "external_key", "required": ["external_key", "po_external_key", "received_quantity"]},
    "QualityStatus": {"entity": "inspection_result", "key": "external_key", "required": ["external_key", "receipt_external_key", "status", "inspected_quantity", "accepted_quantity"]},
    "Invoices": {"entity": "supplier_invoice", "key": "external_key", "required": ["external_key", "supplier_external_key", "po_external_key", "invoice_number", "invoice_date", "subtotal", "tax_amount", "total_amount"]},
    "InvoiceLines": {"entity": "supplier_invoice_line", "key": "external_key", "required": ["external_key", "invoice_external_key", "po_line_external_key", "quantity", "uom", "unit_price", "line_total"]},
    "PaymentStatus": {"entity": "invoice_payment_status", "key": "invoice_external_key", "required": ["invoice_external_key", "status"]},
}
COMMIT_SUPPORTED = frozenset({"company", "plant", "user", "reporting_line", "item", "supplier", "item_specification", "supplier_contact", "supplier_item_capability", "purchase_requirement", "po_draft", "po_draft_line", "store_receipt", "inspection_result", "supplier_invoice", "supplier_invoice_line", "invoice_payment_status"})
ALLOWED_TRANSFORMS = frozenset({"strip", "uppercase", "lowercase"})


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def payload_hash(payload: dict[str, Any]) -> str:
    normalized = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(normalized.encode()).hexdigest()


def safe_spreadsheet_value(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def connector_manifest() -> dict[str, Any]:
    return {
        "connector_type": "excel_csv", "manifest_version": "1.0",
        "operating_modes": ["import_only", "export_package", "managed_round_trip"],
        "capabilities": ["read_materials", "read_suppliers", "read_requirements", "read_pos", "read_receipts", "read_quality", "read_invoices", "read_payment_status", "export_po", "export_reconciliation"],
        "formats": ["xlsx", "csv", "zip_csv"], "write_strategy": "new_version_only",
        "executes_formulas_or_macros": False,
    }


def _clean_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")


def _rows_from_csv(content: bytes) -> list[tuple[str, list[dict[str, Any]]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(422, "CSV must use UTF-8 encoding") from exc
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(422, "CSV has no header row")
    rows: list[dict[str, Any]] = []
    for source_row in reader:
        row: dict[str, Any] = {}
        for source_header, value in source_row.items():
            if source_header is None:
                continue
            header = _clean_header(source_header)
            if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
                row[header] = None
                row[f"__formula_{header}"] = True
            else:
                row[header] = value
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise HTTPException(422, f"CSV exceeds the {MAX_ROWS}-row limit")
    return [("Items", rows)]


def _rows_from_xlsx(content: bytes) -> list[tuple[str, list[dict[str, Any]]]]:
    if not content.startswith(b"PK"):
        raise HTTPException(422, "XLSX signature is invalid")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise HTTPException(422, "Workbook is encrypted, malformed, or unsupported") from exc
    if len(workbook.sheetnames) > MAX_SHEETS:
        raise HTTPException(422, f"Workbook exceeds the {MAX_SHEETS}-sheet limit")
    result: list[tuple[str, list[dict[str, Any]]]] = []
    total = 0
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=False)
        header_cells = next(iterator, ())
        headers = [_clean_header(cell.value) for cell in header_cells]
        if not any(headers):
            continue
        rows: list[dict[str, Any]] = []
        for row in iterator:
            total += 1
            if total > MAX_ROWS:
                raise HTTPException(422, f"Workbook exceeds the {MAX_ROWS}-row limit")
            values: dict[str, Any] = {}
            for index, cell in enumerate(row[:len(headers)]):
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    values[headers[index]] = None
                    values[f"__formula_{headers[index]}"] = True
                else:
                    values[headers[index]] = cell.value
            if any(value not in (None, "") for key, value in values.items() if not key.startswith("__")):
                rows.append(values)
        result.append((sheet.title, rows))
    return result


def parse_workbook(filename: str, content: bytes) -> list[tuple[str, list[dict[str, Any]]]]:
    lower = filename.casefold()
    if lower.endswith(".csv"):
        return _rows_from_csv(content)
    if lower.endswith(".xlsx"):
        return _rows_from_xlsx(content)
    if lower.endswith(".zip"):
        try:
            archive = zipfile.ZipFile(io.BytesIO(content))
        except zipfile.BadZipFile as exc:
            raise HTTPException(422, "ZIP archive is invalid") from exc
        members = [item for item in archive.infolist() if not item.is_dir()]
        if len(members) > MAX_ARCHIVE_FILES or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
            raise HTTPException(422, "ZIP archive exceeds safe extraction limits")
        if any(".." in item.filename.split("/") or not item.filename.casefold().endswith(".csv") for item in members):
            raise HTTPException(422, "ZIP may contain only safely named CSV files")
        return [(item.filename.rsplit("/", 1)[-1][:-4], _rows_from_csv(archive.read(item))[0][1]) for item in members]
    raise HTTPException(422, "Use XLSX, CSV, or a ZIP containing CSV files")


def discover(filename: str, content: bytes) -> dict[str, Any]:
    sheets = parse_workbook(filename, content)
    return {"filename": filename, "source_hash": hashlib.sha256(content).hexdigest(), "sheets": [
        {"name": name, "row_count": len(rows), "headers": sorted({key for row in rows[:20] for key in row if not key.startswith("__")}), "sample": rows[:3], "recognized": name in SHEET_SCHEMAS}
        for name, rows in sheets
    ]}


def validate_mapping_contract(mappings: dict[str, dict[str, str]], transforms: dict[str, dict[str, str]]) -> None:
    optional = {"business_number", "title", "status", "quality_score", "delivery_score", "approval_status", "external_version", "observed_at", "paid_amount", "currency", "reference", "short_quantity", "excess_quantity", "damaged_quantity", "exception_types", "observed_item_code", "certificate_status", "exception_notes", "production_impact", "rejected_quantity", "held_quantity", "defect_codes", "inspection_notes"}
    for sheet, field_map in mappings.items():
        schema = SHEET_SCHEMAS.get(sheet)
        if schema is None:
            raise HTTPException(422, f"Unknown canonical sheet: {sheet}")
        unknown = set(field_map) - set(schema["required"]) - optional
        if unknown:
            raise HTTPException(422, f"Unknown fields for {sheet}: {', '.join(sorted(unknown))}")
        missing = set(schema["required"]) - set(field_map)
        if missing:
            raise HTTPException(422, f"Map required fields for {sheet}: {', '.join(sorted(missing))}")
    for sheet, field_transforms in transforms.items():
        if sheet not in mappings:
            raise HTTPException(422, f"Transforms require a mapping for {sheet}")
        unsupported = set(field_transforms.values()) - ALLOWED_TRANSFORMS
        if unsupported:
            raise HTTPException(422, f"Unsupported transforms: {', '.join(sorted(unsupported))}")


def _apply_mapping(sheet: str, row: dict[str, Any], profile: models.IntegrationMappingProfile | None) -> dict[str, Any]:
    if profile is None or sheet not in (profile.mappings or {}):
        return row
    mapped = {canonical: row.get(_clean_header(source)) for canonical, source in profile.mappings[sheet].items()}
    for field, transform in (profile.transforms or {}).get(sheet, {}).items():
        value = mapped.get(field)
        if not isinstance(value, str):
            continue
        if transform == "strip": mapped[field] = value.strip()
        elif transform == "uppercase": mapped[field] = value.strip().upper()
        elif transform == "lowercase": mapped[field] = value.strip().lower()
    return mapped


ENTITY_MODELS = {
    "company": models.Company, "plant": models.Plant, "user": models.User,
    "reporting_line": models.ReportingLine,
    "item": models.Item, "supplier": models.Supplier,
    "item_specification": models.ItemSpecification, "supplier_contact": models.SupplierContact,
    "supplier_item_capability": models.SupplierItemCapability,
    "purchase_requirement": models.PurchaseRequirement, "po_draft": models.PODraft,
    "po_draft_line": models.PODraftLine, "store_receipt": models.StoreReceipt,
    "inspection_result": models.InspectionResult, "supplier_invoice": models.SupplierInvoice,
    "supplier_invoice_line": models.SupplierInvoiceLine,
    "invoice_payment_status": models.InvoicePaymentStatus,
}


def external_state(db: Session, connection_id: str, entity_type: str, external_key: str | None) -> models.IntegrationExternalRecordState | None:
    if not external_key:
        return None
    return db.query(models.IntegrationExternalRecordState).filter_by(
        connection_id=connection_id, external_entity_type=entity_type, external_key=external_key,
    ).first()


def external_entity(db: Session, connection_id: str, entity_type: str, external_key: str | None):
    state = external_state(db, connection_id, entity_type, external_key)
    model = ENTITY_MODELS.get(entity_type)
    return db.get(model, state.canonical_entity_id) if state and state.canonical_entity_id and model else None


def preview_import(db: Session, user: models.User, connection: models.IntegrationConnection, filename: str, content: bytes, *, document_id: str | None = None, mapping_profile: models.IntegrationMappingProfile | None = None) -> models.IntegrationImportBatch:
    parsed = parse_workbook(filename, content)
    correlation = f"IMP-{secrets.token_hex(8).upper()}"
    discovery = discover(filename, content)
    batch = models.IntegrationImportBatch(
        tenant_id=user.tenant_id, plant_id=user.plant_id, connection_id=connection.id,
        document_id=document_id, mapping_profile_id=mapping_profile.id if mapping_profile else None,
        source_filename=filename, source_hash=discovery["source_hash"],
        status="previewed", mode="dry_run", discovery=discovery, summary={}, correlation_id=correlation,
    )
    db.add(batch)
    db.flush()
    counts = {"insert": 0, "update": 0, "unchanged": 0, "conflict": 0, "error": 0, "ignored": 0}
    seen_external_keys: set[tuple[str, str]] = set()
    for sheet_name, rows in parsed:
        schema = SHEET_SCHEMAS.get(sheet_name)
        for row_number, row in enumerate(rows, start=2):
            row = _apply_mapping(sheet_name, row, mapping_profile)
            messages: list[dict[str, Any]] = []
            if schema is None:
                action, external_key, row_hash = "ignored", None, payload_hash(row)
            else:
                external_key = str(row.get(schema["key"]) or "").strip() or None
                for field in schema["required"]:
                    if row.get(field) in (None, ""):
                        messages.append({"field": field, "code": "required", "message": f"{field.replace('_', ' ').title()} is required"})
                formula_fields = [key.removeprefix("__formula_") for key, value in row.items() if key.startswith("__formula_") and value]
                messages.extend({"field": field, "code": "formula_rejected", "message": "Formula values are not imported; provide a literal value"} for field in formula_fields)
                identity = (schema["entity"], external_key or "")
                if external_key and identity in seen_external_keys:
                    messages.append({"field": schema["key"], "code": "duplicate_external_key", "message": "External key is duplicated in this import batch"})
                elif external_key:
                    seen_external_keys.add(identity)
                row_hash = payload_hash({key: value for key, value in row.items() if not key.startswith("__")})
                state = external_state(db, connection.id, schema["entity"], external_key)
                if messages:
                    action = "error"
                elif schema["entity"] not in COMMIT_SUPPORTED:
                    action = "error"
                    messages.append({"code": "adapter_not_available", "message": f"{sheet_name} can be inspected but its canonical import adapter is not available yet"})
                elif state is None:
                    action = "insert"
                elif state.external_payload_hash == row_hash:
                    action = "unchanged"
                else:
                    model = ENTITY_MODELS.get(schema["entity"])
                    local = db.get(model, state.canonical_entity_id) if model and state.canonical_entity_id else None
                    action = "conflict" if local is not None and state.canonical_version != getattr(local, "version", 1) else "update"
                    if action == "conflict":
                        messages.append({"code": "stale_external_data", "message": "The platform record changed after the last sync; review before overwriting"})
            counts[action] += 1
            db.add(models.IntegrationImportRowResult(
                tenant_id=user.tenant_id, plant_id=user.plant_id, batch_id=batch.id,
                sheet_name=sheet_name, row_number=row_number, external_key=external_key,
                payload_hash=row_hash, action=action, canonical_entity_type=schema["entity"] if schema else None,
                validation_messages=messages, source_cells={key: value for key, value in row.items() if not key.startswith("__")},
                mapping_version=str(mapping_profile.profile_version if mapping_profile else 1), correlation_id=correlation,
            ))
    batch.summary = counts
    return batch


def commit_import(db: Session, user: models.User, batch: models.IntegrationImportBatch) -> models.IntegrationImportBatch:
    if batch.status not in {"previewed", "partially_completed"}:
        raise HTTPException(409, "Only a previewed import can be approved")
    rows = db.query(models.IntegrationImportRowResult).filter_by(batch_id=batch.id).all()
    entity_order = {"company": -4, "plant": -3, "user": -2, "reporting_line": -1, "item": 0, "supplier": 1, "item_specification": 2, "supplier_contact": 2, "supplier_item_capability": 2, "purchase_requirement": 3, "po_draft": 4, "po_draft_line": 5, "store_receipt": 6, "inspection_result": 7, "supplier_invoice": 8, "supplier_invoice_line": 9, "invoice_payment_status": 10}
    rows.sort(key=lambda row: (entity_order.get(row.canonical_entity_type or "", 99), row.sheet_name, row.row_number))
    batch.status, batch.mode, batch.approved_by_user_id, batch.approved_at = "running", "commit", user.id, utcnow()
    for result in rows:
        if result.action not in {"insert", "update"}:
            continue
        row = result.source_cells
        if result.canonical_entity_type == "company":
            entity = external_entity(db, batch.connection_id, "company", result.external_key)
            if entity is None:
                entity = db.query(models.Company).filter_by(tenant_id=user.tenant_id, erp_code=result.external_key).first()
            if entity is None:
                entity = models.Company(id=workflows.next_id(db, models.Company, "COM"), tenant_id=user.tenant_id, name=str(row["name"]), erp_code=result.external_key)
                db.add(entity)
            else:
                entity.name = str(row["name"])
        elif result.canonical_entity_type == "plant":
            company = external_entity(db, batch.connection_id, "company", str(row.get("organization_key") or ""))
            if company is None:
                result.action = "error"; result.validation_messages = [{"field": "organization_key", "code": "reference_not_found", "message": "Import the referenced organization before its plant"}]; continue
            entity = external_entity(db, batch.connection_id, "plant", result.external_key)
            if entity is None:
                entity = db.query(models.Plant).filter_by(tenant_id=user.tenant_id, erp_location_code=str(row["code"])).first()
            # A newly-created workspace starts with one deliberately minimal plant so
            # the owner has a valid scope before master data is connected.  The first
            # authoritative plant in an initial ERP import hydrates that placeholder;
            # creating a second plant here would strand imported users outside the
            # owner's active scope and make the workspace unusable after onboarding.
            tenant = db.get(models.Tenant, user.tenant_id)
            imported_plant_exists = db.query(models.IntegrationExternalRecordState).filter_by(
                connection_id=batch.connection_id,
                external_entity_type="plant",
            ).first()
            if (
                entity is None
                and tenant is not None
                and tenant.workspace_kind == "fresh"
                and tenant.onboarding_status == "needs_master_data"
                and imported_plant_exists is None
            ):
                entity = db.get(models.Plant, user.plant_id)
            if entity is None:
                entity = models.Plant(id=workflows.next_id(db, models.Plant, "PLANT"), tenant_id=user.tenant_id, company_id=company.id, name=str(row["name"]), erp_location_code=str(row["code"]))
                db.add(entity)
            else:
                entity.company_id, entity.name, entity.erp_location_code = company.id, str(row["name"]), str(row["code"])
            if not db.query(models.UserPlantAccess).filter_by(tenant_id=user.tenant_id, user_id=user.id, plant_id=entity.id).first():
                db.add(models.UserPlantAccess(id=workflows.next_id(db, models.UserPlantAccess, "UPA"), tenant_id=user.tenant_id, user_id=user.id, plant_id=entity.id, is_default=entity.id == user.plant_id))
        elif result.canonical_entity_type == "user":
            from app.core.security import hash_password
            role = str(row["role"]).strip().casefold().replace(" ", "_")
            if db.get(models.RoleDefinition, role) is None:
                result.action = "error"; result.validation_messages = [{"field": "role", "code": "invalid_value", "message": "Role is not configured in this workspace"}]; continue
            plant = external_entity(db, batch.connection_id, "plant", str(row.get("plant_key") or ""))
            if plant is None:
                result.action = "error"; result.validation_messages = [{"field": "plant_key", "code": "reference_not_found", "message": "Import the referenced plant before its users"}]; continue
            email = str(row["email"]).strip().casefold()
            entity = external_entity(db, batch.connection_id, "user", result.external_key)
            if entity is None:
                entity = db.query(models.User).filter_by(tenant_id=user.tenant_id, email=email).first()
            if entity is None:
                account = db.query(models.Account).filter_by(email=email).first()
                if account is not None:
                    result.action = "error"; result.validation_messages = [{"field": "email", "code": "email_in_another_workspace", "message": "Use the invitation flow to add an existing account to this workspace"}]; continue
                password_hash = hash_password(secrets.token_urlsafe(32))
                account = models.Account(id=workflows.next_id(db, models.Account, "ACCT"), email=email, name=str(row["name"]), password_hash=password_hash, status="invited")
                department = db.query(models.Department).filter_by(tenant_id=user.tenant_id, plant_id=plant.id, name="Imported procurement").first()
                if department is None:
                    department = models.Department(id=workflows.next_id(db, models.Department, "DEP"), tenant_id=user.tenant_id, plant_id=plant.id, name="Imported procurement"); db.add(department); db.flush()
                entity = models.User(id=workflows.next_id(db, models.User, "USR"), account_id=account.id, tenant_id=user.tenant_id, plant_id=plant.id, department_id=department.id, name=str(row["name"]), email=email, role=role, password_hash=password_hash)
                db.add_all([account, entity]); db.flush()
                membership = models.WorkspaceMembership(id=workflows.next_id(db, models.WorkspaceMembership, "MEM"), account_id=account.id, tenant_id=user.tenant_id, user_id=entity.id, default_plant_id=plant.id, plant_ids=[plant.id], department_id=department.id, role=role, permissions=[], status="active")
                db.add_all([
                    membership,
                    models.UserPlantAccess(id=workflows.next_id(db, models.UserPlantAccess, "UPA"), tenant_id=user.tenant_id, user_id=entity.id, plant_id=plant.id, is_default=True),
                    models.AgentProfile(id=workflows.next_id(db, models.AgentProfile, "AGT"), tenant_id=user.tenant_id, plant_id=plant.id, user_id=entity.id, membership_id=membership.id, role=role, display_name=f"{str(row['name']).split(' ')[0]}'s {role.replace('_', ' ').title()} Agent", allowed_actions=[], blocked_actions=["approve", "publish", "post_to_erp", "update_inventory", "change_master_data"], enabled=False),
                ])
            else:
                entity.name, entity.role, entity.plant_id = str(row["name"]), role, plant.id
        elif result.canonical_entity_type == "reporting_line":
            employee = external_entity(db, batch.connection_id, "user", str(row.get("employee_key") or ""))
            manager = external_entity(db, batch.connection_id, "user", str(row.get("manager_key") or ""))
            if employee is None or manager is None or employee.id == manager.id:
                result.action = "error"; result.validation_messages = [{"code": "invalid_reporting_reference", "message": "Import distinct employee and manager users before their reporting line"}]; continue
            entity = external_entity(db, batch.connection_id, "reporting_line", result.external_key)
            if entity is None:
                entity = db.query(models.ReportingLine).filter_by(tenant_id=user.tenant_id, report_user_id=employee.id).first()
            if entity is None:
                entity = models.ReportingLine(id=workflows.next_id(db, models.ReportingLine, "RPT"), tenant_id=user.tenant_id, manager_user_id=manager.id, report_user_id=employee.id); db.add(entity)
            else:
                entity.manager_user_id = manager.id
            employee.manager_id = manager.id
            employee_membership = db.query(models.WorkspaceMembership).filter_by(tenant_id=user.tenant_id, user_id=employee.id).first()
            manager_membership = db.query(models.WorkspaceMembership).filter_by(tenant_id=user.tenant_id, user_id=manager.id).first()
            if employee_membership and manager_membership:
                employee_membership.manager_membership_id = manager_membership.id
        elif result.canonical_entity_type == "item":
            entity = db.query(models.Item).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id, erp_item_code=result.external_key).first()
            if entity is None:
                entity = models.Item(tenant_id=user.tenant_id, plant_id=user.plant_id, code=str(row["code"]), name=str(row["name"]), uom_id=str(row["uom"]), erp_item_code=result.external_key)
                db.add(entity)
            else:
                entity.code, entity.name, entity.uom_id = str(row["code"]), str(row["name"]), str(row["uom"])
        elif result.canonical_entity_type == "supplier":
            entity = db.query(models.Supplier).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id, erp_vendor_id=result.external_key).first()
            if entity is None:
                entity = models.Supplier(tenant_id=user.tenant_id, plant_id=user.plant_id, name=str(row["name"]), status=str(row.get("status") or "conditional"), quality_score=int(row.get("quality_score") or 0), delivery_score=int(row.get("delivery_score") or 0), erp_vendor_id=result.external_key)
                db.add(entity)
            else:
                entity.name, entity.status = str(row["name"]), str(row.get("status") or entity.status)
        elif result.canonical_entity_type == "item_specification":
            item = external_entity(db, batch.connection_id, "item", str(row.get("item_external_key") or ""))
            if item is None:
                result.action = "error"; result.validation_messages = [{"field": "item_external_key", "code": "reference_not_found", "message": "Import the referenced item before its specification"}]; continue
            entity = external_entity(db, batch.connection_id, "item_specification", result.external_key)
            certificates = [value.strip() for value in str(row.get("required_certificates") or "").split(";") if value.strip()]
            inspection_required = str(row.get("inspection_required", "true")).casefold() not in {"false", "0", "no"}
            if entity is None:
                entity = models.ItemSpecification(tenant_id=user.tenant_id, plant_id=user.plant_id, item_id=item.id, description=str(row["description"]), required_certificates=certificates, inspection_required=inspection_required); db.add(entity)
            else:
                entity.description, entity.required_certificates, entity.inspection_required = str(row["description"]), certificates, inspection_required
        elif result.canonical_entity_type == "supplier_contact":
            supplier = external_entity(db, batch.connection_id, "supplier", str(row.get("supplier_external_key") or ""))
            if supplier is None:
                result.action = "error"; result.validation_messages = [{"field": "supplier_external_key", "code": "reference_not_found", "message": "Import the referenced supplier before its contact"}]; continue
            entity = external_entity(db, batch.connection_id, "supplier_contact", result.external_key)
            if entity is None:
                entity = models.SupplierContact(tenant_id=user.tenant_id, plant_id=user.plant_id, supplier_id=supplier.id, name=str(row["name"]), email=str(row["email"]), phone=str(row.get("phone") or "")); db.add(entity)
            else:
                entity.name, entity.email, entity.phone = str(row["name"]), str(row["email"]), str(row.get("phone") or "")
        elif result.canonical_entity_type == "supplier_item_capability":
            supplier = external_entity(db, batch.connection_id, "supplier", str(row.get("supplier_external_key") or ""))
            item = external_entity(db, batch.connection_id, "item", str(row.get("item_external_key") or ""))
            if supplier is None or item is None:
                result.action = "error"; result.validation_messages = [{"code": "reference_not_found", "message": "Import the referenced supplier and item before capability"}]; continue
            entity = external_entity(db, batch.connection_id, "supplier_item_capability", result.external_key)
            approved = str(row.get("approved", "true")).casefold() not in {"false", "0", "no"}
            if entity is None:
                entity = models.SupplierItemCapability(tenant_id=user.tenant_id, plant_id=user.plant_id, supplier_id=supplier.id, item_id=item.id, approved=approved, notes=str(row.get("notes") or "")); db.add(entity)
            else:
                entity.approved, entity.notes = approved, str(row.get("notes") or "")
        elif result.canonical_entity_type == "purchase_requirement":
            item = db.query(models.Item).filter_by(
                tenant_id=user.tenant_id, plant_id=user.plant_id, code=str(row["item_code"]),
            ).first()
            if item is None:
                result.action = "error"
                result.validation_messages = [{"field": "item_code", "code": "reference_not_found", "message": "Import Items first or use an approved canonical item code"}]
                continue
            entity = workflows.create_requirement(
                db, user, item.id, float(row["quantity"]), str(row["need_by_date"]),
                uom=str(row["uom"]), title=str(row.get("title") or f"{item.name} requirement"),
                reason=str(row["reason"]), source="excel",
            )
            if row.get("business_number"):
                entity.business_number = str(row["business_number"])
        elif result.canonical_entity_type == "po_draft":
            supplier = external_entity(db, batch.connection_id, "supplier", str(row.get("supplier_external_key") or ""))
            if supplier is None:
                result.action = "error"; result.validation_messages = [{"field": "supplier_external_key", "code": "reference_not_found", "message": "Import the referenced supplier before this purchase order"}]; continue
            entity = external_entity(db, batch.connection_id, "po_draft", result.external_key)
            status = str(row.get("status") or "posted")
            if status not in {"approved_pending_outbox", "posting", "simulated_posted", "posted", "failed", "rejected"}:
                result.action = "error"; result.validation_messages = [{"field": "status", "code": "invalid_value", "message": "Purchase-order status is not supported"}]; continue
            if entity is None:
                entity = models.PODraft(
                    tenant_id=user.tenant_id, plant_id=user.plant_id, award_id=None,
                    supplier_id=supplier.id, status=status,
                    oracle_mapping={"external_owned": True, "external_key": result.external_key},
                    currency=str(row.get("currency") or "INR"), payment_terms=str(row.get("payment_terms") or ""),
                    commercial_terms={}, business_number=str(row.get("business_number") or result.external_key),
                ); db.add(entity)
            else:
                entity.status = status
        elif result.canonical_entity_type == "po_draft_line":
            po = external_entity(db, batch.connection_id, "po_draft", str(row.get("po_external_key") or ""))
            item = external_entity(db, batch.connection_id, "item", str(row.get("item_external_key") or ""))
            if po is None or item is None:
                result.action = "error"; result.validation_messages = [{"code": "reference_not_found", "message": "Import the referenced purchase order and item before its line"}]; continue
            entity = external_entity(db, batch.connection_id, "po_draft_line", result.external_key)
            if entity is None:
                entity = models.PODraftLine(tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id, item_id=item.id, quantity=float(row["quantity"]), uom=str(row["uom"]), unit_price=float(row["unit_price"]), gst_rate=float(row.get("gst_rate") or 0), need_by_date=str(row["need_by_date"]), inspection_required=str(row.get("inspection_required", "true")).casefold() not in {"false", "0", "no"}); db.add(entity)
            else:
                entity.quantity, entity.unit_price, entity.need_by_date = float(row["quantity"]), float(row["unit_price"]), str(row["need_by_date"])
        elif result.canonical_entity_type == "store_receipt":
            po = external_entity(db, batch.connection_id, "po_draft", str(row.get("po_external_key") or ""))
            if po is None:
                result.action = "error"; result.validation_messages = [{"field": "po_external_key", "code": "reference_not_found", "message": "Import the referenced purchase order before its receipt"}]; continue
            entity = external_entity(db, batch.connection_id, "store_receipt", result.external_key)
            exceptions = row.get("exception_types") if isinstance(row.get("exception_types"), list) else [value.strip() for value in str(row.get("exception_types") or "").split(",") if value.strip()]
            production_impact = str(row.get("production_impact") or "false").casefold() in {"true", "1", "yes"}
            values = dict(status=str(row.get("status") or "received"), received_quantity=float(row["received_quantity"]), short_quantity=float(row.get("short_quantity") or 0), excess_quantity=float(row.get("excess_quantity") or 0), damaged_quantity=float(row.get("damaged_quantity") or 0), exception_types=exceptions, observed_item_code=str(row.get("observed_item_code") or ""), certificate_status=str(row.get("certificate_status") or "received"), exception_notes=str(row.get("exception_notes") or ""), production_impact=production_impact)
            if entity is None:
                entity = models.StoreReceipt(tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id, business_number=str(row.get("business_number") or result.external_key), **values); db.add(entity)
            else:
                for field, value in values.items(): setattr(entity, field, value)
        elif result.canonical_entity_type == "inspection_result":
            receipt = external_entity(db, batch.connection_id, "store_receipt", str(row.get("receipt_external_key") or ""))
            if receipt is None:
                result.action = "error"; result.validation_messages = [{"field": "receipt_external_key", "code": "reference_not_found", "message": "Import the referenced receipt before quality status"}]; continue
            entity = external_entity(db, batch.connection_id, "inspection_result", result.external_key)
            values = (float(row["inspected_quantity"]), float(row["accepted_quantity"]), float(row.get("rejected_quantity") or 0), float(row.get("held_quantity") or 0))
            if sum(values[1:]) > values[0]:
                result.action = "error"; result.validation_messages = [{"code": "invalid_disposition", "message": "Accepted, rejected, and held quantities cannot exceed inspected quantity"}]; continue
            defects = row.get("defect_codes") if isinstance(row.get("defect_codes"), list) else [value.strip() for value in str(row.get("defect_codes") or "").split(",") if value.strip()]
            quality_values = dict(status=str(row.get("status") or "completed"), inspected_quantity=values[0], accepted_quantity=values[1], rejected_quantity=values[2], held_quantity=values[3], certificate_status=str(row.get("certificate_status") or "verified"), defect_codes=defects, inspection_notes=str(row.get("inspection_notes") or ""), production_impact=str(row.get("production_impact") or "false").casefold() in {"true", "1", "yes"})
            if entity is None:
                entity = models.InspectionResult(tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=receipt.po_draft_id, receipt_id=receipt.id, business_number=str(row.get("business_number") or result.external_key), **quality_values); db.add(entity)
            else:
                for field, value in quality_values.items(): setattr(entity, field, value)
        elif result.canonical_entity_type == "supplier_invoice":
            supplier = external_entity(db, batch.connection_id, "supplier", str(row.get("supplier_external_key") or ""))
            po = external_entity(db, batch.connection_id, "po_draft", str(row.get("po_external_key") or ""))
            if supplier is None or po is None or po.supplier_id != supplier.id:
                result.action = "error"; result.validation_messages = [{"code": "reference_not_found", "message": "Import matching supplier and purchase-order references before the invoice"}]; continue
            subtotal, tax, total = float(row["subtotal"]), float(row["tax_amount"]), float(row["total_amount"])
            if abs(total - subtotal - tax) > 0.01:
                result.action = "error"; result.validation_messages = [{"code": "invalid_total", "message": "Invoice total must equal subtotal plus tax"}]; continue
            entity = external_entity(db, batch.connection_id, "supplier_invoice", result.external_key)
            if entity is None:
                duplicate = workflows.user_scope_query(db, models.SupplierInvoice, user).filter_by(supplier_id=supplier.id, invoice_number=str(row["invoice_number"])).first()
                if duplicate:
                    result.action = "error"; result.validation_messages = [{"code": "duplicate_invoice", "message": "This supplier invoice number already exists"}]; continue
                entity = models.SupplierInvoice(tenant_id=user.tenant_id, plant_id=user.plant_id, supplier_id=supplier.id, po_draft_id=po.id, invoice_number=str(row["invoice_number"]), invoice_date=str(row["invoice_date"]), currency=str(row.get("currency") or po.currency), subtotal=subtotal, tax_amount=tax, total_amount=total, status=str(row.get("status") or "captured"), source="excel", business_number=str(row.get("business_number") or result.external_key)); db.add(entity)
            else:
                entity.subtotal, entity.tax_amount, entity.total_amount = subtotal, tax, total
        elif result.canonical_entity_type == "supplier_invoice_line":
            invoice = external_entity(db, batch.connection_id, "supplier_invoice", str(row.get("invoice_external_key") or ""))
            po_line = external_entity(db, batch.connection_id, "po_draft_line", str(row.get("po_line_external_key") or ""))
            if invoice is None or po_line is None or invoice.po_draft_id != po_line.po_draft_id:
                result.action = "error"; result.validation_messages = [{"code": "reference_not_found", "message": "Import matching invoice and purchase-order line references first"}]; continue
            entity = external_entity(db, batch.connection_id, "supplier_invoice_line", result.external_key)
            if entity is None:
                entity = models.SupplierInvoiceLine(tenant_id=user.tenant_id, plant_id=user.plant_id, invoice_id=invoice.id, po_line_id=po_line.id, item_id=po_line.item_id, description=str(row.get("description") or ""), quantity=float(row["quantity"]), uom=str(row["uom"]), unit_price=float(row["unit_price"]), tax_amount=float(row.get("tax_amount") or 0), line_total=float(row["line_total"])); db.add(entity)
            else:
                entity.quantity, entity.unit_price, entity.tax_amount, entity.line_total = float(row["quantity"]), float(row["unit_price"]), float(row.get("tax_amount") or 0), float(row["line_total"])
        elif result.canonical_entity_type == "invoice_payment_status":
            invoice = external_entity(db, batch.connection_id, "supplier_invoice", str(row.get("invoice_external_key") or ""))
            if invoice is None:
                result.action = "error"; result.validation_messages = [{"field": "invoice_external_key", "code": "reference_not_found", "message": "Import the referenced invoice before payment status"}]; continue
            status = str(row.get("status") or "unknown").strip().casefold().replace(" ", "_")
            allowed_statuses = {"not_released", "scheduled", "processing", "paid", "partially_paid", "on_hold", "rejected", "unknown"}
            if status not in allowed_statuses:
                result.action = "error"; result.validation_messages = [{"field": "status", "code": "invalid_value", "message": "Payment status is not recognized"}]; continue
            version = str(row.get("external_version") or "")
            entity = workflows.user_scope_query(db, models.InvoicePaymentStatus, user).filter_by(
                connection_id=batch.connection_id, external_key=result.external_key,
                external_version=version, source_payload_hash=result.payload_hash,
            ).first()
            if entity is None:
                observed_raw = row.get("observed_at")
                try:
                    observed_at = datetime.fromisoformat(str(observed_raw)) if observed_raw else utcnow()
                except ValueError:
                    result.action = "error"; result.validation_messages = [{"field": "observed_at", "code": "invalid_date", "message": "Observed time must use ISO date/time format"}]; continue
                entity = models.InvoicePaymentStatus(
                    tenant_id=user.tenant_id, plant_id=user.plant_id, invoice_id=invoice.id,
                    connection_id=batch.connection_id, external_key=result.external_key,
                    status=status, external_version=version, observed_at=observed_at,
                    source_payload_hash=result.payload_hash,
                    evidence={"paid_amount": row.get("paid_amount"), "currency": row.get("currency"), "reference": row.get("reference"), "source": "excel_read_only"},
                ); db.add(entity)
        else:
            continue
        db.flush()
        state = db.query(models.IntegrationExternalRecordState).filter_by(connection_id=batch.connection_id, external_entity_type=result.canonical_entity_type, external_key=result.external_key).first()
        if state is None:
            state = models.IntegrationExternalRecordState(tenant_id=user.tenant_id, plant_id=user.plant_id, connection_id=batch.connection_id, external_entity_type=result.canonical_entity_type, external_key=result.external_key, external_payload_hash=result.payload_hash, ownership="external_owned", sync_status="synced")
            db.add(state)
        state.canonical_entity_type, state.canonical_entity_id = result.canonical_entity_type, entity.id
        state.external_payload_hash, state.canonical_version, state.last_synced_at = result.payload_hash, getattr(entity, "version", 1), utcnow()
        result.canonical_entity_id, result.action = entity.id, "committed"
        workflows.create_audit(
            db, user.tenant_id, user.plant_id, user.name, "excel_import.row_committed",
            result.canonical_entity_type or "external_record", entity.id, actor_user_id=user.id,
            meta={"batch_id": batch.id, "sheet": result.sheet_name, "row": result.row_number, "external_key": result.external_key, "payload_hash": result.payload_hash},
        )
        # Sessions disable autoflush; dependent rows in this same batch must be
        # able to resolve the external record state just persisted above.
        db.flush()
    db.flush()
    batch.status = "completed_with_errors" if any(row.action in {"error", "conflict"} for row in rows) else "completed"
    imported_types = {row.canonical_entity_type for row in rows if row.canonical_entity_type}
    seen_keys = {(row.canonical_entity_type, row.external_key) for row in rows if row.external_key}
    missing_states = db.query(models.IntegrationExternalRecordState).filter(
        models.IntegrationExternalRecordState.connection_id == batch.connection_id,
        models.IntegrationExternalRecordState.external_entity_type.in_(imported_types or {""}),
        models.IntegrationExternalRecordState.tombstoned.is_(False),
    ).all()
    missing = [state.external_key for state in missing_states if (state.external_entity_type, state.external_key) not in seen_keys]
    batch.summary = {
        **batch.summary, "committed": sum(row.action == "committed" for row in rows),
        "missing_external_records": len(missing),
        "deletion_policy": "retain_and_report",
        "tombstones_applied": 0,
    }
    return batch


def export_exchange_workbook(db: Session, user: models.User) -> bytes:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["GenuineGigs Procurement Exchange", "Schema version", "1.0"])
    readme.append(["Safety", "Values are data only; formulas, macros and external links are never executed."])
    for name, headers, rows in (
        ("Items", ["external_key", "code", "name", "uom", "approval_status"], db.query(models.Item).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all()),
        ("Suppliers", ["external_key", "name", "status", "quality_score", "delivery_score"], db.query(models.Supplier).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all()),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            values = [row.erp_item_code, row.code, row.name, row.uom_id, "approved"] if name == "Items" else [row.erp_vendor_id, row.name, row.status, row.quality_score, row.delivery_score]
            sheet.append([safe_spreadsheet_value(value) for value in values])
    states = db.query(models.IntegrationExternalRecordState).filter_by(
        tenant_id=user.tenant_id, plant_id=user.plant_id,
    ).all()
    external_keys = {(state.external_entity_type, state.canonical_entity_id): state.external_key for state in states}

    def key(entity_type: str, entity_id: str | None) -> str:
        return external_keys.get((entity_type, entity_id), "")

    organization_sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    companies = [row for row in db.query(models.Company).filter_by(tenant_id=user.tenant_id).all() if key("company", row.id)]
    organization_sheets.append(("Organizations", ["external_key", "name"], [[key("company", row.id), row.name] for row in companies]))
    plants = [row for row in db.query(models.Plant).filter_by(tenant_id=user.tenant_id).all() if key("plant", row.id)]
    organization_sheets.append(("Plants", ["external_key", "organization_key", "code", "name"], [[key("plant", row.id), key("company", row.company_id), row.erp_location_code, row.name] for row in plants]))
    imported_users = [row for row in db.query(models.User).filter_by(tenant_id=user.tenant_id).all() if key("user", row.id)]
    organization_sheets.append(("Users", ["external_key", "name", "email", "role", "plant_key"], [[key("user", row.id), row.name, row.email, row.role, key("plant", row.plant_id)] for row in imported_users]))
    reporting_lines = [row for row in db.query(models.ReportingLine).filter_by(tenant_id=user.tenant_id).all() if key("reporting_line", row.id)]
    organization_sheets.append(("ReportingLines", ["employee_key", "manager_key"], [[key("user", row.report_user_id), key("user", row.manager_user_id)] for row in reporting_lines]))
    for name, headers, rows in organization_sheets:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for values in rows:
            sheet.append([safe_spreadsheet_value(value) for value in values])

    transaction_sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    all_pos = db.query(models.PODraft).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all()
    pos = [po for po in all_pos if key("po_draft", po.id)]
    transaction_sheets.append(("OpenPOs", ["external_key", "business_number", "supplier_external_key", "status", "currency", "payment_terms", "canonical_version"], [[key("po_draft", po.id), po.business_number, key("supplier", po.supplier_id), po.status, po.currency, po.payment_terms, po.version] for po in pos]))
    all_po_lines = db.query(models.PODraftLine).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all()
    po_lines = [line for line in all_po_lines if key("po_draft_line", line.id)]
    transaction_sheets.append(("OpenPOLines", ["external_key", "po_external_key", "item_external_key", "quantity", "uom", "unit_price", "gst_rate", "need_by_date"], [[key("po_draft_line", line.id), key("po_draft", line.po_draft_id), key("item", line.item_id), line.quantity, line.uom, line.unit_price, line.gst_rate, line.need_by_date] for line in po_lines]))
    receipts = [row for row in db.query(models.StoreReceipt).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all() if key("store_receipt", row.id)]
    transaction_sheets.append(("Receipts", ["external_key", "business_number", "po_external_key", "status", "received_quantity", "short_quantity", "excess_quantity", "damaged_quantity", "exception_types", "observed_item_code", "certificate_status", "exception_notes", "production_impact"], [[key("store_receipt", row.id), row.business_number, key("po_draft", row.po_draft_id), row.status, row.received_quantity, row.short_quantity, row.excess_quantity, row.damaged_quantity, ",".join(row.exception_types or []), row.observed_item_code, row.certificate_status, row.exception_notes, row.production_impact] for row in receipts]))
    inspections = [row for row in db.query(models.InspectionResult).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all() if key("inspection_result", row.id)]
    transaction_sheets.append(("QualityStatus", ["external_key", "business_number", "receipt_external_key", "status", "inspected_quantity", "accepted_quantity", "rejected_quantity", "held_quantity", "certificate_status", "defect_codes", "inspection_notes", "production_impact"], [[key("inspection_result", row.id), row.business_number, key("store_receipt", row.receipt_id), row.status, row.inspected_quantity, row.accepted_quantity, row.rejected_quantity, row.held_quantity, row.certificate_status, ",".join(row.defect_codes or []), row.inspection_notes, row.production_impact] for row in inspections]))
    invoices = [row for row in db.query(models.SupplierInvoice).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all() if key("supplier_invoice", row.id)]
    transaction_sheets.append(("Invoices", ["external_key", "business_number", "supplier_external_key", "po_external_key", "invoice_number", "invoice_date", "subtotal", "tax_amount", "total_amount", "currency", "status"], [[key("supplier_invoice", row.id), row.business_number, key("supplier", row.supplier_id), key("po_draft", row.po_draft_id), row.invoice_number, row.invoice_date, row.subtotal, row.tax_amount, row.total_amount, row.currency, row.status] for row in invoices]))
    invoice_lines = [row for row in db.query(models.SupplierInvoiceLine).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).all() if key("supplier_invoice_line", row.id)]
    transaction_sheets.append(("InvoiceLines", ["external_key", "invoice_external_key", "po_line_external_key", "quantity", "uom", "unit_price", "tax_amount", "line_total"], [[key("supplier_invoice_line", row.id), key("supplier_invoice", row.invoice_id), key("po_draft_line", row.po_line_id), row.quantity, row.uom, row.unit_price, row.tax_amount, row.line_total] for row in invoice_lines]))
    payment_statuses = db.query(models.InvoicePaymentStatus).filter_by(tenant_id=user.tenant_id, plant_id=user.plant_id).order_by(models.InvoicePaymentStatus.observed_at).all()
    transaction_sheets.append(("PaymentStatus", ["invoice_external_key", "status", "external_version", "observed_at", "paid_amount", "currency", "reference"], [[key("supplier_invoice", row.invoice_id), row.status, row.external_version, row.observed_at.isoformat(), row.evidence.get("paid_amount"), row.evidence.get("currency"), row.evidence.get("reference")] for row in payment_statuses]))
    for name, headers, rows in transaction_sheets:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for values in rows:
            sheet.append([safe_spreadsheet_value(value) for value in values])
    po_export = workbook.create_sheet("POExport")
    po_export.append(["export_key", "business_number", "supplier_id", "status", "currency", "payment_terms", "canonical_version", "source"])
    for po in all_pos:
        po_export.append([safe_spreadsheet_value(key("po_draft", po.id) or f"GG-PO:{po.business_number or po.id}"), po.business_number, po.supplier_id, po.status, po.currency, po.payment_terms, po.version, "external_round_trip" if key("po_draft", po.id) else "genuinegigs"])
    reconciliation = workbook.create_sheet("Reconciliation")
    reconciliation.append(["entity_type", "external_key", "canonical_id", "sync_status", "external_payload_hash", "canonical_version", "last_synced_at"])
    for state in states:
        reconciliation.append([state.external_entity_type, safe_spreadsheet_value(state.external_key), state.canonical_entity_id, state.sync_status, state.external_payload_hash, state.canonical_version, state.last_synced_at.isoformat() if state.last_synced_at else ""])
    sync = workbook.create_sheet("SyncLog")
    sync.append(["generated_at", "tenant", "plant", "mode", "external_records", "reconciled_records"])
    sync.append([utcnow().isoformat(), user.tenant_id, user.plant_id, "new_version", len(states), sum(state.sync_status == "synced" for state in states)])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def verify_exchange_export(content: bytes, expected_states: list[models.IntegrationExternalRecordState]) -> dict[str, Any]:
    """Read the generated artifact back and prove every external identity survived."""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        rows = list(workbook["Reconciliation"].iter_rows(min_row=2, values_only=True))
    except Exception as exc:
        return {"status": "failed", "error": "generated_workbook_unreadable", "verified": 0, "missing": []}
    exported = {(str(row[0]), str(row[1])) for row in rows if row and row[0] and row[1]}
    expected = {(state.external_entity_type, state.external_key) for state in expected_states}
    missing = sorted(f"{entity_type}:{external_key}" for entity_type, external_key in expected - exported)
    return {"status": "reconciled" if not missing else "mismatch", "verified": len(expected & exported), "expected": len(expected), "missing": missing}


def import_reconciliation_report(db: Session, user: models.User, batch: models.IntegrationImportBatch) -> bytes:
    """Return a spreadsheet-safe, row-level CSV report for customer reconciliation."""
    rows = db.query(models.IntegrationImportRowResult).filter_by(
        tenant_id=user.tenant_id, plant_id=user.plant_id, batch_id=batch.id,
    ).order_by(models.IntegrationImportRowResult.sheet_name, models.IntegrationImportRowResult.row_number).all()
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([
        "source_file", "source_hash", "batch_status", "sheet", "row", "external_key",
        "action", "canonical_entity_type", "canonical_entity_id", "messages", "mapping_version",
    ])
    for row in rows:
        messages = "; ".join(
            f"{message.get('code', 'review_required')}: {message.get('message', 'Review required')}"
            for message in row.validation_messages
        )
        writer.writerow([
            safe_spreadsheet_value(batch.source_filename), batch.source_hash, batch.status,
            safe_spreadsheet_value(row.sheet_name), row.row_number,
            safe_spreadsheet_value(row.external_key or ""), row.action,
            row.canonical_entity_type or "", row.canonical_entity_id or "",
            safe_spreadsheet_value(messages), row.mapping_version,
        ])
    return output.getvalue().encode("utf-8-sig")


def create_exchange_export(db: Session, user: models.User) -> dict[str, Any]:
    """Create and verify one immutable workbook version for manual or agent callers."""
    from app.storage import object_storage

    workflows.ensure_default_connections(db, user)
    db.flush()
    connection = workflows.user_scope_query(db, models.IntegrationConnection, user).filter(
        models.IntegrationConnection.provider == "local",
    ).first()
    if connection is None:
        raise HTTPException(409, "Excel connector is not configured")
    content = export_exchange_workbook(db, user)
    settings = workflows.get_settings()
    version = workflows.user_scope_query(db, models.IntegrationExportBatch, user).count() + 1
    filename = f"genuinegigs-procurement-exchange-v{version}.xlsx"
    key = f"{user.tenant_id}/{user.plant_id}/excel-exports/{secrets.token_hex(12)}-{filename}"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    object_storage.put_file(settings.object_storage_evidence_bucket, key, io.BytesIO(content), content_type)
    checksum = hashlib.sha256(content).hexdigest()
    document = models.Document(
        tenant_id=user.tenant_id, plant_id=user.plant_id, filename=filename, content_type=content_type,
        size_bytes=len(content), storage_key=key, storage_bucket=settings.object_storage_evidence_bucket,
        checksum_sha256=checksum, status="available", validation_errors=[],
        linked_entity_type="integration_export", uploaded_by_user_id=user.id,
    )
    db.add(document)
    db.flush()
    batch = models.IntegrationExportBatch(
        tenant_id=user.tenant_id, plant_id=user.plant_id, connection_id=connection.id,
        output_document_id=document.id, export_type="procurement_exchange", status="completed",
        output_version=version, manifest=connector_manifest(),
        summary={"filename": filename, "size_bytes": len(content)},
        correlation_id=f"EXP-{secrets.token_hex(8).upper()}",
    )
    db.add(batch)
    db.flush()
    states = workflows.user_scope_query(db, models.IntegrationExternalRecordState, user).all()
    verification = verify_exchange_export(content, states)
    batch.status = "completed" if verification["status"] == "reconciled" else "verification_failed"
    batch.summary = {**batch.summary, "read_after_write": verification}
    db.add(models.ReconciliationResult(
        tenant_id=user.tenant_id, plant_id=user.plant_id, provider="excel_csv",
        entity_type="integration_export_batch", entity_id=batch.id, status=verification["status"],
        local_hash=checksum, external_hash=checksum,
        differences={"missing": verification.get("missing", []), "verified": verification.get("verified", 0)},
    ))
    workflows.create_audit(
        db, user.tenant_id, user.plant_id, user.name, "excel_export.created",
        "integration_export_batch", batch.id, actor_user_id=user.id,
        meta={"output_document_id": document.id, "version": version},
    )
    return {
        "export_batch_id": batch.id, "document_id": document.id, "filename": filename,
        "status": batch.status, "reconciliation": verification,
        "download_url": object_storage.presigned_get(document.storage_bucket, document.storage_key, 300),
    }
