from datetime import timedelta
from io import BytesIO
import hashlib
import asyncio

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app import agent_service, documents
from app.agent_schemas import AgentMessageRequest, IntentEnvelope, ThreadCreateRequest
from app.db import models
from app.db.seed import reset_workspace_database
from app.db.session import SessionLocal
from app.domains import workflows
from app.main import supplier_po_context, supplier_po_document, supplier_rfq_context, supplier_rfq_document


@pytest.fixture(autouse=True)
def reset_workspace() -> None:
    with SessionLocal() as db:
        reset_workspace_database(db); db.commit()


def setup_portal(db):
    user = db.query(models.User).filter_by(email="purchase.manager@genuinegigs.local").one()
    award = db.query(models.AwardDecision).first()
    quote_line = db.query(models.QuoteLine).filter_by(quote_id=award.quote_id).first()
    po = models.PODraft(
        business_number="PO-PORTAL-001", tenant_id=user.tenant_id, plant_id=user.plant_id,
        award_id=award.id, supplier_id=award.supplier_id, status="posted",
        oracle_mapping={}, currency="INR", payment_terms="30 days", commercial_terms={},
    )
    db.add(po); db.flush()
    db.add(models.PODraftLine(
        tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id,
        item_id=quote_line.item_id, quantity=100, uom="KG", unit_price=100,
        gst_rate=18, need_by_date="2026-08-10", inspection_required=True,
    ))
    raw_token = "supplier-delivery-token"
    token = models.SupplierPortalToken(
        tenant_id=user.tenant_id, plant_id=user.plant_id, rfq_id="",
        supplier_id=po.supplier_id, token_hash=hashlib.sha256(raw_token.encode()).hexdigest(),
        status="active", expires_at=workflows.utcnow() + timedelta(days=2),
        purpose="po_acknowledgement", entity_id=po.id,
    )
    db.add(token); db.flush()
    return user, po, raw_token, token


def setup_rfq_portal(db):
    user = db.query(models.User).filter_by(email="purchase.exec@genuinegigs.local").one()
    rfq = db.query(models.RFQ).first()
    supplier_id = (rfq.supplier_ids or [db.query(models.Supplier).first().id])[0]
    if supplier_id not in (rfq.supplier_ids or []): rfq.supplier_ids = [supplier_id]
    invitation = db.query(models.RFQSupplierInvitation).filter_by(rfq_id=rfq.id, supplier_id=supplier_id).first()
    if not invitation:
        invitation = models.RFQSupplierInvitation(
            tenant_id=rfq.tenant_id, plant_id=rfq.plant_id, rfq_id=rfq.id,
            supplier_id=supplier_id, status="sent",
        ); db.add(invitation)
    else: invitation.status = "sent"
    raw_token = "supplier-rfq-token"
    db.add(models.SupplierPortalToken(
        tenant_id=rfq.tenant_id, plant_id=rfq.plant_id, rfq_id=rfq.id,
        supplier_id=supplier_id, token_hash=hashlib.sha256(raw_token.encode()).hexdigest(),
        status="active", expires_at=workflows.utcnow() + timedelta(days=2),
        purpose="rfq_quote", entity_id=rfq.id,
    )); db.flush()
    return user, rfq, invitation, raw_token


def test_supplier_acknowledgement_delivery_and_clarification_are_scoped_and_idempotent() -> None:
    with SessionLocal() as db:
        user, po, token, portal = setup_portal(db)
        acknowledgement = workflows.receive_supplier_acknowledgement(db, token, "accepted", 100, "2026-08-08", "Confirmed")
        db.flush()
        assert workflows.receive_supplier_acknowledgement(db, token, "accepted", 100, "2026-08-08", "Confirmed").id == acknowledgement.id
        assert portal.status == "active" and portal.used_at is not None
        with pytest.raises(HTTPException) as changed:
            workflows.receive_supplier_acknowledgement(db, token, "rejected", 0, "2026-08-08", "Changed mind")
        assert changed.value.status_code == 409

        first = workflows.receive_supplier_delivery_notice(db, token, 60, "2026-08-05", "DISP-001", "MH12AB1234")
        db.flush()
        assert first.source == "supplier_portal" and first.status == "supplier_submitted"
        assert workflows.receive_supplier_delivery_notice(db, token, 60, "2026-08-05", "DISP-001", "MH12AB1234").id == first.id
        update = workflows.receive_supplier_delivery_update(
            db, token, first.id, "TRACK-UPDATE-1", "2026-08-07", "Carrier route was disrupted by flooding",
        )
        db.flush()
        assert update.update_type == "delay" and first.expected_delivery == "2026-08-07"
        assert workflows.receive_supplier_delivery_update(
            db, token, first.id, "TRACK-UPDATE-1", "2026-08-07", "Carrier route was disrupted by flooding",
        ).id == update.id
        with pytest.raises(HTTPException) as changed_update:
            workflows.receive_supplier_delivery_update(
                db, token, first.id, "TRACK-UPDATE-1", "2026-08-06", "Different commitment",
            )
        assert changed_update.value.status_code == 409
        assert db.query(models.Task).filter_by(semantic_key=f"supplier_delivery_update:{update.id}:coordinate").one().severity == "critical"
        assert db.query(models.Notification).filter_by(linked_entity_id=update.id).count() == 2
        context = supplier_po_context(token, db)
        first_context = next(row for row in context["deliveries"] if row["id"] == first.id)
        assert first_context["updates"][0]["reason"] == "Carrier route was disrupted by flooding"
        second = workflows.receive_supplier_delivery_notice(db, token, 40, "2026-08-08", "DISP-002")
        db.flush()
        with pytest.raises(HTTPException) as excess:
            workflows.receive_supplier_delivery_notice(db, token, 1, "2026-08-09", "DISP-003")
        assert excess.value.status_code == 422
        assert db.query(models.Task).filter_by(semantic_key=f"supplier_asn:{first.id}:receive").one().owner_role == "store_manager"
        assert db.query(models.Notification).filter_by(dedupe_key=f"supplier_asn:{second.id}:receive").one()

        message = workflows.receive_supplier_clarification(db, token, "PORTAL-MSG-1", "Delivery question", "Can Gate receive after 6 PM?")
        db.flush()
        assert message.supplier_id == po.supplier_id and message.direction == "inbound"
        assert workflows.receive_supplier_clarification(db, token, "PORTAL-MSG-1", "Delivery question", "Can Gate receive after 6 PM?").id == message.id
        assert db.query(models.Task).filter_by(semantic_key=f"supplier_clarification:{message.id}:respond").one()


def test_supplier_dispatch_document_is_quarantined_validated_and_linked(monkeypatch) -> None:
    workbook = Workbook(); workbook.active.append(["Dispatch", "DISP-001"])
    content = BytesIO(); workbook.save(content)
    with SessionLocal() as db:
        _user, _po, token, _portal = setup_portal(db)
        workflows.receive_supplier_acknowledgement(db, token, "accepted", 100, "2026-08-08", "Confirmed")
        asn = workflows.receive_supplier_delivery_notice(db, token, 100, "2026-08-08", "DISP-001")
        db.flush()
        monkeypatch.setattr(documents.object_storage, "put_file", lambda *_args: None)
        monkeypatch.setattr(documents.object_storage, "get_bytes", lambda *_args: content.getvalue())
        monkeypatch.setattr(documents.object_storage, "copy", lambda *_args: None)
        class AsyncUpload:
            filename = "dispatch.xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            def __init__(self, value: bytes): self.stream = BytesIO(value)
            async def read(self, size: int = -1) -> bytes: return self.stream.read(size)
        upload = AsyncUpload(content.getvalue())
        document, job = asyncio.run(documents.create_supplier_portal_upload(db, asn, upload, "supplier_dispatch_document", asn.id))
        workflows.attach_supplier_delivery_document(db, token, asn.id, document); db.flush()
        assert document.status == "quarantined" and document.id in asn.supplier_document_ids
        documents.process_document_job(db, job.id); db.flush()
        assert document.status == "available" and job.status == "completed"
        assert db.query(models.QuoteExtractionRun).filter_by(document_id=document.id).count() == 0


def test_supplier_invoice_upload_reuses_canonical_extraction_and_is_idempotent(monkeypatch) -> None:
    workbook = Workbook()
    workbook.active.append(["Invoice Number", "INV-SUP-001"])
    workbook.active.append(["Invoice Date", "2026-08-09"])
    workbook.active.append(["PO Number", "PO-PORTAL-001"])
    workbook.active.append(["Quantity", 100])
    workbook.active.append(["Unit Price", 100])
    workbook.active.append(["Grand Total", 11800])
    content = BytesIO(); workbook.save(content); raw = content.getvalue()
    with SessionLocal() as db:
        _user, po, token, _portal = setup_portal(db)
        workflows.receive_supplier_acknowledgement(db, token, "accepted", 100, "2026-08-08", "Confirmed")
        db.flush()
        monkeypatch.setattr(documents.object_storage, "put_file", lambda *_args: None)
        monkeypatch.setattr(documents.object_storage, "get_bytes", lambda *_args: raw)
        monkeypatch.setattr(documents.object_storage, "copy", lambda *_args: None)

        class AsyncUpload:
            filename = "invoice.xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            def __init__(self, value: bytes): self.stream = BytesIO(value)
            async def read(self, size: int = -1) -> bytes: return self.stream.read(size)

        document, job = asyncio.run(documents.create_supplier_portal_upload(
            db, po, AsyncUpload(raw), "supplier_invoice_draft", po.id,
        ))
        same_document, same_job = asyncio.run(documents.create_supplier_portal_upload(
            db, po, AsyncUpload(raw), "supplier_invoice_draft", po.id,
        ))
        assert same_document.id == document.id and same_job.id == job.id
        assert job.job_type == "validate_extract"
        db.flush()
        documents.process_document_job(db, job.id); db.flush()
        extraction = db.query(models.InvoiceExtractionRun).filter_by(document_id=document.id).one()
        assert extraction.po_draft_id == po.id and extraction.status == "needs_review"
        assert document.status == "validated_needs_review"
        assert db.query(models.SupplierInvoice).filter_by(document_id=document.id).count() == 0
        assert db.query(models.Task).filter_by(semantic_key=f"invoice_extraction:{extraction.id}:review").one()


def test_supplier_can_download_only_the_final_document_for_token_po(monkeypatch) -> None:
    pdf = b"%PDF-1.4\ncontrolled purchase order"
    with SessionLocal() as db:
        user, po, token, _portal = setup_portal(db)
        document = models.Document(
            tenant_id=po.tenant_id, plant_id=po.plant_id, filename="PO-PORTAL-001.pdf",
            content_type="application/pdf", size_bytes=len(pdf), storage_key="po/final.pdf",
            storage_bucket="evidence", checksum_sha256=hashlib.sha256(pdf).hexdigest(),
            status="generated_final", validation_errors=[], linked_entity_type="po", linked_entity_id=po.id,
        )
        db.add(document); db.flush()
        db.add(models.GeneratedArtifact(
            tenant_id=po.tenant_id, plant_id=po.plant_id, artifact_type="po", entity_type="po",
            entity_id=po.id, artifact_version=1, document_id=document.id, status="final",
            source_snapshot_hash="snapshot", generated_by_user_id=user.id,
        )); db.flush()
        monkeypatch.setattr(documents.object_storage, "get_bytes", lambda bucket, key: pdf)
        context = supplier_po_context(token, db)
        assert context["purchase_order"]["document_available"] is True
        response = supplier_po_document(token, db)
        assert response.body == pdf
        assert response.headers["cache-control"] == "private, no-store"


def test_supplier_rfq_context_and_no_bid_are_scoped_and_idempotent() -> None:
    with SessionLocal() as db:
        _user, rfq, invitation, token = setup_rfq_portal(db)
        context = supplier_rfq_context(token, db)
        assert context["supplier_request"]["business_number"] == rfq.business_number
        assert context["response_status"] == "sent" and context["lines"]
        result = workflows.receive_supplier_no_bid(db, token, "Current production capacity is fully committed")
        db.flush()
        assert result.id == invitation.id and invitation.status == "no_bid"
        assert workflows.receive_supplier_no_bid(db, token, "Current production capacity is fully committed").id == invitation.id
        with pytest.raises(HTTPException) as changed:
            workflows.receive_supplier_no_bid(db, token, "Different reason")
        assert changed.value.status_code == 409
        assert db.query(models.SupplierChannelMessage).filter_by(external_message_id=f"no-bid:{rfq.id}:{invitation.supplier_id}").one()
        assert db.query(models.Notification).filter_by(dedupe_key=f"no-bid:{rfq.id}:{invitation.supplier_id}").one()


def test_supplier_rfq_document_quote_status_and_evidence_are_canonical(monkeypatch) -> None:
    pdf = b"%PDF-1.4\ncontrolled supplier request"
    workbook = Workbook(); workbook.active.append(["Quote Number", "Q-SUP-001"])
    evidence_stream = BytesIO(); workbook.save(evidence_stream); evidence = evidence_stream.getvalue()
    with SessionLocal() as db:
        user, rfq, invitation, token = setup_rfq_portal(db)
        rfq.status = "published"
        document = models.Document(
            tenant_id=rfq.tenant_id, plant_id=rfq.plant_id, filename=f"{rfq.business_number}.pdf",
            content_type="application/pdf", size_bytes=len(pdf), storage_key="rfq/final.pdf",
            storage_bucket="evidence", checksum_sha256=hashlib.sha256(pdf).hexdigest(),
            status="generated_final", validation_errors=[], linked_entity_type="rfq", linked_entity_id=rfq.id,
        )
        db.add(document); db.flush()
        db.add(models.GeneratedArtifact(
            tenant_id=rfq.tenant_id, plant_id=rfq.plant_id, artifact_type="rfq", entity_type="rfq",
            entity_id=rfq.id, artifact_version=1, document_id=document.id, status="final",
            source_snapshot_hash="rfq-snapshot", generated_by_user_id=user.id,
        )); db.flush()
        monkeypatch.setattr(documents.object_storage, "get_bytes", lambda bucket, key: pdf if key == "rfq/final.pdf" else evidence)
        monkeypatch.setattr(documents.object_storage, "put_file", lambda *_args: None)
        monkeypatch.setattr(documents.object_storage, "copy", lambda *_args: None)
        response = supplier_rfq_document(token, db)
        assert response.body == pdf and response.headers["cache-control"] == "private, no-store"

        line = db.query(models.RFQLine).filter_by(rfq_id=rfq.id).first()
        quote = workflows.receive_supplier_quote(db, token, {
            "quote_number": "Q-SUP-001", "currency": "INR",
            "lines": [{"rfq_line_id": line.id, "quantity": line.quantity, "unit_price": 125, "promised_date": line.need_by_date}],
        })
        db.flush()
        assert invitation.status == "responded" and quote.verification_status == "needs_review"

        class AsyncUpload:
            filename = "quotation-evidence.xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            def __init__(self, value: bytes): self.stream = BytesIO(value)
            async def read(self, size: int = -1) -> bytes: return self.stream.read(size)

        uploaded, job = asyncio.run(documents.create_supplier_portal_upload(
            db, quote, AsyncUpload(evidence), "supplier_quote", quote.id,
        ))
        replay, replay_job = asyncio.run(documents.create_supplier_portal_upload(
            db, quote, AsyncUpload(evidence), "supplier_quote", quote.id,
        ))
        assert replay.id == uploaded.id and replay_job.id == job.id
        db.flush(); documents.process_document_job(db, job.id); db.flush()
        extraction = db.query(models.QuoteExtractionRun).filter_by(document_id=uploaded.id).one()
        assert extraction.quote_id == quote.id and uploaded.linked_entity_id == quote.id


def test_store_agent_lists_only_supplier_delivery_notices(monkeypatch) -> None:
    with SessionLocal() as db:
        _manager, _po, token, _portal = setup_portal(db)
        workflows.receive_supplier_acknowledgement(db, token, "accepted", 100, "2026-08-08", "Confirmed")
        asn = workflows.receive_supplier_delivery_notice(db, token, 100, "2026-08-08", "DISP-AGENT")
        workflows.receive_supplier_delivery_update(
            db, token, asn.id, "AGENT-UPDATE-1", "2026-08-10", "Vehicle breakdown",
        )
        db.flush()
        store = db.query(models.User).filter_by(email="store.manager@genuinegigs.local").one()
        thread = agent_service.create_thread(db, store, ThreadCreateRequest(thread_type="personal", title="Supplier delivery")); db.flush()
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (
            IntentEnvelope(intent="list_supplier_deliveries", requested_outcome="Show supplier deliveries", entities={}, confidence=1),
            "groq", "test-model", None,
        ))
        run, message = agent_service.run_message(db, store, thread.id, AgentMessageRequest(content="What has the supplier dispatched?"))
        assert run.intent == "list_supplier_deliveries"
        assert "supplier delivery" in message.content.casefold()
        assert any(asn.business_number in item for block in message.content_blocks for item in block.get("items", []))
        assert any("Vehicle breakdown" in item for block in message.content_blocks for item in block.get("items", []))
        assert db.query(models.AgentActionReceipt).filter_by(run_id=run.id).count() == 0
