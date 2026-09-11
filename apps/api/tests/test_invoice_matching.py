from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app import agent_service, documents, invoice_service
from app.agent_schemas import AgentMessageRequest, IntentEnvelope, ThreadCreateRequest
from app.db import models
from app.db.seed import reset_workspace_database
from app.db.session import SessionLocal


@pytest.fixture(autouse=True)
def reset_workspace() -> None:
    with SessionLocal() as db:
        reset_workspace_database(db); db.commit()


def setup_order(db):
    user = db.query(models.User).filter_by(email="purchase.manager@genuinegigs.local").one()
    award = db.query(models.AwardDecision).first()
    quote_line = db.query(models.QuoteLine).filter_by(quote_id=award.quote_id).first()
    po = models.PODraft(tenant_id=user.tenant_id, plant_id=user.plant_id, award_id=award.id, supplier_id=award.supplier_id, status="posted", oracle_mapping={}, currency="INR", payment_terms="30 days", commercial_terms={})
    db.add(po); db.flush()
    line = models.PODraftLine(tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id, item_id=quote_line.item_id, quantity=10, uom="KG", unit_price=100, gst_rate=18, need_by_date="2026-08-01", inspection_required=True)
    db.add(line); db.flush()
    return user, po, line


def invoice_payload(po, line, *, invoice_number="INV-001", quantity=10, total=1180):
    return {"supplier_id": po.supplier_id, "po_draft_id": po.id, "invoice_number": invoice_number, "invoice_date": "2026-07-19", "currency": "INR", "total_amount": total, "lines": [{"po_line_id": line.id, "quantity": quantity, "uom": "KG", "unit_price": 100, "line_total": quantity * 100, "tax_amount": quantity * 18}]}


def test_three_way_match_and_finance_handoff_never_authorize_payment() -> None:
    with SessionLocal() as db:
        user, po, line = setup_order(db)
        db.add(models.StoreReceipt(tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id, status="received", received_quantity=10, short_quantity=0, excess_quantity=0, damaged_quantity=0))
        invoice = invoice_service.create_invoice(db, user, invoice_payload(po, line)); db.flush()
        match = invoice_service.match_invoice(db, user, invoice.id); db.flush()
        assert (match.match_type, match.result) == ("three_way", "matched")
        handoff = invoice_service.prepare_finance_handoff(db, user, match.id); db.flush()
        assert handoff.status == "prepared_for_finance_review"
        assert handoff.payload["payment_release"] == "not_authorized"
        assert invoice.status == "finance_handoff_prepared"


def test_match_exception_blocks_finance_and_duplicate_invoice_is_rejected() -> None:
    with SessionLocal() as db:
        user, po, line = setup_order(db)
        invoice = invoice_service.create_invoice(db, user, invoice_payload(po, line, quantity=12, total=1416)); db.flush()
        match = invoice_service.match_invoice(db, user, invoice.id); db.flush()
        assert match.result == "exception"
        assert "invoice_quantity_exceeds_order" in match.variances["exceptions"]
        task = db.query(models.Task).filter_by(semantic_key=f"invoice_match:{invoice.id}:exception").one()
        assert task.status == "blocked" and task.owner_role == "purchase_manager"
        assert db.query(models.Notification).filter_by(dedupe_key=f"invoice_match:{invoice.id}:exception").one()
        with pytest.raises(HTTPException) as blocked:
            invoice_service.prepare_finance_handoff(db, user, match.id)
        assert blocked.value.status_code == 409
        with pytest.raises(HTTPException) as duplicate:
            invoice_service.create_invoice(db, user, invoice_payload(po, line, quantity=12, total=1416))
        assert duplicate.value.status_code == 409

        invoice_line = db.query(models.SupplierInvoiceLine).filter_by(invoice_id=invoice.id).one()
        invoice_line.quantity = 10; invoice_line.line_total = 1000; invoice_line.tax_amount = 180
        invoice.subtotal = 1000; invoice.tax_amount = 180; invoice.total_amount = 1180
        corrected = invoice_service.match_invoice(db, user, invoice.id); db.flush()
        assert corrected.result == "matched" and match.status == "superseded"
        assert task.status == "completed"


def test_agent_matches_invoice_then_requires_confirmation_for_finance_handoff(monkeypatch) -> None:
    with SessionLocal() as db:
        user, po, line = setup_order(db)
        db.add(models.StoreReceipt(tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id, status="received", received_quantity=10, short_quantity=0, excess_quantity=0, damaged_quantity=0))
        invoice = invoice_service.create_invoice(db, user, invoice_payload(po, line)); db.flush()
        thread = agent_service.create_thread(db, user, ThreadCreateRequest(thread_type="personal", title="Invoice agent")); db.flush()
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (IntentEnvelope(intent="match_supplier_invoice", requested_outcome="Match invoice", entities={"invoice_id": invoice.id}, confidence=1), "groq", "test-model", None))
        _run, message = agent_service.run_message(db, user, thread.id, AgentMessageRequest(content="Match this invoice"))
        assert "match passed" in message.content.casefold()
        match = db.query(models.InvoiceMatchResult).filter_by(invoice_id=invoice.id).one()

        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (IntentEnvelope(intent="prepare_finance_handoff_proposal", requested_outcome="Prepare finance review", entities={"match_id": match.id}, confidence=1), "groq", "test-model", None))
        _run, proposal_message = agent_service.run_message(db, user, thread.id, AgentMessageRequest(content="Prepare it for finance review"))
        proposal_id = next(block["record"]["proposal_id"] for block in proposal_message.content_blocks if block["type"] == "action_proposal")
        assert db.query(models.FinanceHandoff).count() == 0
        agent_service.confirm_proposal(db, user, proposal_id); db.flush()
        handoff = db.query(models.FinanceHandoff).filter_by(match_result_id=match.id).one()
        assert handoff.payload["payment_release"] == "not_authorized"


def test_agent_reports_read_only_external_payment_observation_without_claiming_release(monkeypatch) -> None:
    with SessionLocal() as db:
        user, po, line = setup_order(db)
        invoice = invoice_service.create_invoice(db, user, invoice_payload(po, line)); db.flush()
        connection = models.IntegrationConnection(
            tenant_id=user.tenant_id, plant_id=user.plant_id, provider="excel_csv",
            provider_version="1.0", name="Finance workbook", mode="read_only", status="configured",
            capabilities=["read_payment_status"], enabled_capabilities=["read_payment_status"], writes_enabled=False,
        )
        db.add(connection); db.flush()
        db.add(models.InvoicePaymentStatus(
            tenant_id=user.tenant_id, plant_id=user.plant_id, invoice_id=invoice.id,
            connection_id=connection.id, external_key="EXT-INV-001", status="scheduled",
            external_version="7", observed_at=invoice.created_at, source_payload_hash="a" * 64,
            evidence={"source": "excel_read_only"},
        )); db.flush()
        thread = agent_service.create_thread(db, user, ThreadCreateRequest(thread_type="personal", title="Payment status")); db.flush()
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (
            IntentEnvelope(intent="check_invoice_payment_status", requested_outcome="Check payment status", entities={"invoice_id": invoice.business_number}, confidence=1),
            "provider-fallback", "test-model", None,
        ))
        _run, message = agent_service.run_message(db, user, thread.id, AgentMessageRequest(content=f"Check payment status for {invoice.business_number}"))
        assert "scheduled" in message.content.casefold()
        assert "read-only" in message.content.casefold()
        assert "released" not in message.content.casefold()
        assert db.query(models.AgentActionReceipt).filter_by(run_id=_run.id).count() == 0


def test_uploaded_invoice_is_extracted_then_requires_human_verification(monkeypatch) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Tax Invoice No", "SUP-INV-778"])
    sheet.append(["Invoice Date", "2026-07-22"])
    sheet.append(["PO Number", "PO-TEST-001"])
    sheet.append(["Quantity", 10])
    sheet.append(["UOM", "KG"])
    sheet.append(["Unit Price", 100])
    sheet.append(["Tax Amount", 180])
    sheet.append(["Grand Total", 1180])
    content = BytesIO(); workbook.save(content)

    with SessionLocal() as db:
        user, po, line = setup_order(db)
        document = models.Document(
            tenant_id=user.tenant_id, plant_id=user.plant_id, filename="supplier-invoice.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(content.getvalue()), storage_key="test/invoice.xlsx", storage_bucket="quarantine",
            checksum_sha256="b" * 64, status="quarantined", validation_errors=[],
            linked_entity_type="supplier_invoice_draft", linked_entity_id=po.id,
            uploaded_by_user_id=user.id,
        )
        db.add(document); db.flush()
        job = models.DocumentJob(
            tenant_id=user.tenant_id, plant_id=user.plant_id, document_id=document.id, status="queued",
        )
        db.add(job); db.flush()
        monkeypatch.setattr(documents.object_storage, "get_bytes", lambda *_args: content.getvalue())
        monkeypatch.setattr(documents.object_storage, "copy", lambda *_args: None)
        documents.process_document_job(db, job.id); db.flush()
        extraction = db.query(models.InvoiceExtractionRun).filter_by(document_id=document.id).one()
        assert extraction.status == "needs_review"
        assert extraction.extracted_fields["invoice_number"] == "SUP-INV-778"
        assert extraction.extracted_fields["total_amount"] == 1180
        assert db.query(models.SupplierInvoice).count() == 0
        review_task = db.query(models.Task).filter_by(semantic_key=f"invoice_extraction:{extraction.id}:review").one()
        assert review_task.owner_role == "purchase_manager" and review_task.status == "open"
        assert db.query(models.Notification).filter_by(dedupe_key=f"invoice_extraction:{extraction.id}:review").one()

        payload = invoice_payload(po, line, invoice_number="SUP-INV-778") | {"source": "upload"}
        invoice = invoice_service.verify_invoice_extraction(db, user, extraction.id, payload); db.flush()
        assert invoice.document_id == document.id
        assert invoice.content_hash == document.checksum_sha256
        assert extraction.invoice_id == invoice.id and extraction.status == "verified"
        assert extraction.verified_by_user_id == user.id
        assert review_task.status == "completed"
        assert document.linked_entity_type == "supplier_invoice"
        assert invoice_service.verify_invoice_extraction(db, user, extraction.id, payload).id == invoice.id

        duplicate_payload = invoice_payload(po, line, invoice_number="SUP-INV-779") | {
            "source": "upload", "content_hash": document.checksum_sha256,
        }
        with pytest.raises(HTTPException) as duplicate:
            invoice_service.create_invoice(db, user, duplicate_payload)
        assert duplicate.value.status_code == 409


def test_agent_prepares_validated_invoice_attachment_for_review(monkeypatch) -> None:
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["Tax Invoice No", "AGENT-INV-42"]); sheet.append(["Invoice Date", "2026-07-22"])
    sheet.append(["Quantity", 10]); sheet.append(["Unit Price", 100]); sheet.append(["Grand Total", 1180])
    content = BytesIO(); workbook.save(content)
    with SessionLocal() as db:
        user, po, _line = setup_order(db)
        document = models.Document(
            tenant_id=user.tenant_id, plant_id=user.plant_id, filename="agent-invoice.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(content.getvalue()), storage_key="evidence/agent-invoice.xlsx", storage_bucket="evidence",
            checksum_sha256="c" * 64, status="validated_needs_review", validation_errors=[], uploaded_by_user_id=user.id,
        )
        db.add(document); db.flush()
        thread = agent_service.create_thread(db, user, ThreadCreateRequest(thread_type="personal", title="Invoice attachment")); db.flush()
        monkeypatch.setattr("app.storage.object_storage.get_bytes", lambda *_args: content.getvalue())
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (
            IntentEnvelope(intent="prepare_invoice_from_attachment", requested_outcome="Prepare this invoice", entities={"po_id": po.id}, attachment_ids=[document.id], confidence=1),
            "groq", "test-model", None,
        ))
        run, message = agent_service.run_message(db, user, thread.id, AgentMessageRequest(content="Prepare this supplier invoice", attachment_ids=[document.id]))
        extraction = db.query(models.InvoiceExtractionRun).filter_by(document_id=document.id).one()
        assert run.intent == "prepare_invoice_from_attachment"
        assert extraction.po_draft_id == po.id and extraction.status == "needs_review"
        assert "ready for your review" in message.content
        assert db.query(models.AgentActionReceipt).filter_by(run_id=run.id).count() == 1
