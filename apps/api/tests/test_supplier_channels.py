from pathlib import Path
from io import BytesIO

import pytest
from openpyxl import Workbook

from app import excel_connector
from app.api_schemas import InboundSupplierEmailRequest
from app.db import models
from app.db.seed import reset_workspace_database
from app.db.session import SessionLocal
from app.domains import workflows
from app.integrations import process_inbound_supplier_email
from app.storage import object_storage


@pytest.fixture(autouse=True)
def reset_workspace() -> None:
    with SessionLocal() as db:
        reset_workspace_database(db)
        db.commit()


def _email_context(db):
    fixture = Path(__file__).parent / "fixtures" / "excel_erp" / "genuinegigs_procurement_exchange_v1.xlsx"
    user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
    workflows.ensure_default_connections(db, user); db.flush()
    connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
    batch = excel_connector.preview_import(db, user, connection, fixture.name, fixture.read_bytes()); db.flush()
    excel_connector.commit_import(db, user, batch); db.flush()
    requirement = db.query(models.PurchaseRequirement).filter_by(business_number="REQ-EXT-0001").one()
    rfq = workflows.create_rfq_from_requirement(db, user, requirement.id); db.flush()
    document = models.Document(
        tenant_id=user.tenant_id, plant_id=user.plant_id, filename="supplier-quote.pdf",
        content_type="application/pdf", size_bytes=128, storage_key="tests/supplier-quote.pdf",
        storage_bucket="evidence", checksum_sha256="b" * 64, status="available",
        validation_errors=[], uploaded_by_user_id=user.id,
    )
    db.add(document); db.flush()
    return user, rfq, document


def test_inbound_email_creates_same_canonical_quote_and_is_idempotent() -> None:
    with SessionLocal() as db:
        user, rfq, document = _email_context(db)
        payload = InboundSupplierEmailRequest(
            external_message_id="MAIL-GOLDEN-1", sender_email="supplier1@fixture.invalid",
            rfq_reference=rfq.business_number, subject=f"Quotation for {rfq.business_number}",
            document_ids=[document.id],
        )
        message, results = process_inbound_supplier_email(db, user, payload); db.flush()
        assert message.status == "processed" and len(results) == 1
        quote = db.get(models.SupplierQuote, results[0]["quote_id"])
        assert quote is not None and quote.rfq_id == rfq.id and quote.supplier_id == message.supplier_id
        assert document.linked_entity_type == "supplier_quote" and document.linked_entity_id == quote.id
        repeated, repeated_results = process_inbound_supplier_email(db, user, payload); db.flush()
        assert repeated.id == message.id and repeated_results[0]["quote_id"] == quote.id
        assert db.query(models.SupplierQuote).filter_by(rfq_id=rfq.id).count() == 1


def test_inbound_email_rejects_unmatched_sender_and_cross_workspace_document() -> None:
    with SessionLocal() as db:
        user, rfq, document = _email_context(db)
        with pytest.raises(ValueError, match="supplier contact"):
            process_inbound_supplier_email(db, user, InboundSupplierEmailRequest(
                external_message_id="MAIL-BAD-SENDER", sender_email="attacker@invalid.test",
                rfq_reference=rfq.id, document_ids=[document.id],
            ))
        document.tenant_id = "another-tenant"; db.flush()
        with pytest.raises(ValueError, match="outside this workspace"):
            process_inbound_supplier_email(db, user, InboundSupplierEmailRequest(
                external_message_id="MAIL-BAD-DOC", sender_email="supplier1@fixture.invalid",
                rfq_reference=rfq.id, document_ids=[document.id],
            ))


def test_inbound_invoice_email_creates_review_not_payable_and_is_idempotent(monkeypatch) -> None:
    workbook = Workbook()
    workbook.active.append(["Invoice Number", "INV-EMAIL-001"])
    workbook.active.append(["Invoice Date", "2026-07-24"])
    workbook.active.append(["Quantity", 200])
    workbook.active.append(["Unit Price", 100])
    workbook.active.append(["Grand Total", 23600])
    stream = BytesIO(); workbook.save(stream); content = stream.getvalue()
    with SessionLocal() as db:
        user, _rfq, _quote_document = _email_context(db)
        po = db.query(models.PODraft).filter_by(business_number="PO-EXT-0001").one()
        document = models.Document(
            tenant_id=user.tenant_id, plant_id=user.plant_id, filename="supplier-invoice.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(content), storage_key="tests/supplier-invoice.xlsx",
            storage_bucket="evidence", checksum_sha256="c" * 64, status="available",
            validation_errors=[], uploaded_by_user_id=user.id,
        )
        db.add(document); db.flush()
        monkeypatch.setattr(object_storage, "get_bytes", lambda *_args: content)
        payload = InboundSupplierEmailRequest(
            external_message_id="MAIL-INVOICE-1", sender_email="supplier1@fixture.invalid",
            message_type="invoice", po_reference=po.business_number,
            subject=f"Invoice for {po.business_number}", document_ids=[document.id],
        )
        message, results = process_inbound_supplier_email(db, user, payload); db.flush()
        extraction = db.get(models.InvoiceExtractionRun, results[0]["invoice_extraction_id"])
        assert message.status == "processed" and message.rfq_id is None
        assert extraction is not None and extraction.po_draft_id == po.id and extraction.status == "needs_review"
        assert document.linked_entity_type == "supplier_invoice_draft" and document.linked_entity_id == po.id
        assert db.query(models.SupplierInvoice).filter_by(document_id=document.id).count() == 0
        assert db.query(models.Task).filter_by(semantic_key=f"invoice_extraction:{extraction.id}:review").count() == 1
        repeated, repeated_results = process_inbound_supplier_email(db, user, payload); db.flush()
        assert repeated.id == message.id and repeated_results[0]["invoice_extraction_id"] == extraction.id
        assert db.query(models.InvoiceExtractionRun).filter_by(document_id=document.id).count() == 1


def test_inbound_email_acknowledgement_and_delivery_use_canonical_services() -> None:
    with SessionLocal() as db:
        user, _rfq, delivery_document = _email_context(db)
        po = db.query(models.PODraft).filter_by(business_number="PO-EXT-0001").one()
        ordered = sum(line.quantity for line in db.query(models.PODraftLine).filter_by(po_draft_id=po.id).all())
        acknowledgement_payload = InboundSupplierEmailRequest(
            external_message_id="MAIL-ACK-1", sender_email="supplier1@fixture.invalid",
            message_type="po_acknowledgement", po_reference=po.business_number,
            acknowledgement_status="accepted", confirmed_quantity=ordered,
            confirmed_delivery="2026-08-12", subject=f"Accepted {po.business_number}",
            body_preview="We accept the order and delivery commitment.",
        )
        message, results = process_inbound_supplier_email(db, user, acknowledgement_payload); db.flush()
        acknowledgement = db.get(models.SupplierAcknowledgement, results[0]["acknowledgement_id"])
        assert acknowledgement is not None and acknowledgement.po_draft_id == po.id and acknowledgement.status == "accepted"
        replay, replay_results = process_inbound_supplier_email(db, user, acknowledgement_payload); db.flush()
        assert replay.id == message.id and replay_results[0]["acknowledgement_id"] == acknowledgement.id

        delivery_payload = InboundSupplierEmailRequest(
            external_message_id="MAIL-ASN-1", sender_email="supplier1@fixture.invalid",
            message_type="delivery_notice", po_reference=po.business_number,
            expected_quantity=ordered, expected_delivery="2026-08-12",
            dispatch_reference="EMAIL-DISP-001", vehicle_number="MH12EMAIL",
            subject=f"Dispatch for {po.business_number}", document_ids=[delivery_document.id],
            attachment_purpose="supplier_certificate",
        )
        delivery_message, delivery_results = process_inbound_supplier_email(db, user, delivery_payload); db.flush()
        asn = db.get(models.ASN, delivery_results[0]["asn_id"])
        assert delivery_message.status == "processed"
        assert asn is not None and asn.po_draft_id == po.id and asn.source == "supplier_email"
        assert delivery_document.linked_entity_type == "supplier_certificate"
        assert delivery_document.linked_entity_id == asn.id and delivery_document.id in asn.supplier_document_ids
        assert db.query(models.Task).filter_by(semantic_key=f"supplier_asn:{asn.id}:receive").one().owner_role == "store_manager"
        repeated_delivery, repeated_results = process_inbound_supplier_email(db, user, delivery_payload); db.flush()
        assert repeated_delivery.id == delivery_message.id and repeated_results[0]["asn_id"] == asn.id
        assert db.query(models.ASN).filter_by(po_draft_id=po.id, dispatch_reference="EMAIL-DISP-001").count() == 1
