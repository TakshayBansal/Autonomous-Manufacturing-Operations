"""Regenerate the deterministic Excel external-system fixture pack."""
from pathlib import Path
import json
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).parent

SHEETS = [
    "README", "Organizations", "Plants", "Users", "ReportingLines", "UOMs",
    "Currencies", "TaxCodes", "Items", "ItemSpecifications", "Suppliers",
    "SupplierContacts", "SupplierCapabilities", "InventorySnapshot", "Requirements",
    "RequirementLines", "OpenPOs", "OpenPOLines", "Receipts", "ReceiptLines",
    "QualityStatus", "Invoices", "InvoiceLines", "PaymentStatus", "RFQExport",
    "ComparisonExport", "POExport", "ExceptionExport", "SyncLog",
]


def build(path: Path, *, stale: bool = False, hostile: bool = False) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEETS:
        workbook.create_sheet(name)
    workbook["README"].append(["GenuineGigs Procurement Exchange", "schema_version", "1.0"])
    workbook["Organizations"].append(["external_key", "name"]); workbook["Organizations"].append(["ORG-APEX", "Apex Components"])
    workbook["Plants"].append(["external_key", "organization_key", "code", "name"]); workbook["Plants"].append(["PLANT-PUNE", "ORG-APEX", "PUNE", "Pune Plant"]); workbook["Plants"].append(["PLANT-NASHIK", "ORG-APEX", "NASHIK", "Nashik Plant"])
    workbook["Users"].append(["external_key", "name", "email", "role", "plant_key"])
    roles = ["plant_manager", "purchase_manager", "purchase_executive", "gate_operator", "store_manager", "quality_inspector", "admin"]
    for index, role in enumerate(roles, 1): workbook["Users"].append([f"USER-{index}", role.replace("_", " ").title(), f"{role}@fixture.invalid", role, "PLANT-PUNE"])
    workbook["ReportingLines"].append(["employee_key", "manager_key"]); workbook["ReportingLines"].append(["USER-3", "USER-2"])
    workbook["UOMs"].append(["code", "label"]); workbook["UOMs"].append(["KG", "Kilogram"]); workbook["UOMs"].append(["EA", "Each"])
    workbook["Currencies"].append(["code", "label"]); workbook["Currencies"].append(["INR", "Indian Rupee"])
    workbook["TaxCodes"].append(["code", "rate", "recoverable"]); workbook["TaxCodes"].append(["GST18", 18, True])
    workbook["Items"].append(["external_key", "code", "name", "uom"])
    materials = ["Aluminium Ingot", "Copper Ingot", "Nickel Ingot", "Steel Round Bar", "Bearing Sleeve", "Zinc Ingot", "Brass Rod", "Fastener Kit", "Cutting Oil", "Packaging Crate"]
    for index, name in enumerate(materials, 1): workbook["Items"].append([f"EXT-ITEM-{index:03}", f"MAT-{index:03}", "=2+2" if hostile and index == 1 else name, "KG" if index < 8 else "EA"])
    workbook["ItemSpecifications"].append(["external_key", "item_external_key", "description"]); workbook["ItemSpecifications"].append(["SPEC-1", "EXT-ITEM-001", "Primary aluminium, approved grade"])
    workbook["Suppliers"].append(["external_key", "name", "status", "quality_score", "delivery_score"])
    for index in range(1, 7): workbook["Suppliers"].append([f"EXT-SUP-{index:03}", f"Fixture Supplier {index}", "approved", 80 + index, 78 + index])
    workbook["SupplierContacts"].append(["external_key", "supplier_external_key", "name", "email"]); workbook["SupplierContacts"].append(["CONTACT-1", "EXT-SUP-001", "Supplier Desk", "supplier1@fixture.invalid"])
    workbook["SupplierCapabilities"].append(["external_key", "supplier_external_key", "item_external_key", "approved"]); workbook["SupplierCapabilities"].append(["CAP-1", "EXT-SUP-001", "EXT-ITEM-001", True])
    workbook["InventorySnapshot"].append(["external_key", "item_external_key", "quantity", "uom", "as_of"]); workbook["InventorySnapshot"].append(["INV-1", "EXT-ITEM-001", 120, "KG", "2026-07-19"])
    workbook["Requirements"].append(["external_key", "business_number", "item_code", "quantity", "uom", "need_by_date", "reason"]); workbook["Requirements"].append(["EXT-REQ-001", "REQ-EXT-0001", "MAT-001", 200, "KG", "2026-08-18", "Production raw material"])
    workbook["RequirementLines"].append(["external_key", "requirement_external_key", "item_external_key", "quantity", "uom"]); workbook["RequirementLines"].append(["EXT-REQL-001", "EXT-REQ-001", "EXT-ITEM-001", 200, "KG"])
    workbook["OpenPOs"].append(["external_key", "business_number", "supplier_external_key", "status", "currency", "payment_terms", "external_version"]); workbook["OpenPOs"].append(["EXT-PO-001", "PO-EXT-0001", "EXT-SUP-001", "posted", "INR", "30 days", "0" if stale else "2"])
    workbook["OpenPOLines"].append(["external_key", "po_external_key", "item_external_key", "quantity", "uom", "unit_price", "gst_rate", "need_by_date"]); workbook["OpenPOLines"].append(["EXT-POL-001", "EXT-PO-001", "EXT-ITEM-001", 200, "KG", 100, 18, "2026-08-18"])
    workbook["Receipts"].append(["external_key", "business_number", "po_external_key", "status", "received_quantity", "short_quantity", "excess_quantity", "damaged_quantity"]); workbook["Receipts"].append(["EXT-REC-001", "GRN-EXT-0001", "EXT-PO-001", "received", 180, 20, 0, 0])
    workbook["ReceiptLines"].append(["external_key", "receipt_external_key", "item_external_key", "quantity"]); workbook["ReceiptLines"].append(["EXT-RECL-001", "EXT-REC-001", "EXT-ITEM-001", 180])
    workbook["QualityStatus"].append(["external_key", "business_number", "receipt_external_key", "status", "inspected_quantity", "accepted_quantity", "rejected_quantity", "held_quantity"]); workbook["QualityStatus"].append(["EXT-QA-001", "INS-EXT-0001", "EXT-REC-001", "hold", 180, 170, 0, 10])
    workbook["Invoices"].append(["external_key", "business_number", "supplier_external_key", "po_external_key", "invoice_number", "invoice_date", "subtotal", "tax_amount", "total_amount", "currency"]); workbook["Invoices"].append(["EXT-INV-001", "INV-EXT-0001", "EXT-SUP-001", "EXT-PO-001", "INV-FIX-001", "2026-07-19", 18000, 3240, 21240, "INR"])
    workbook["InvoiceLines"].append(["external_key", "invoice_external_key", "po_line_external_key", "quantity", "uom", "unit_price", "tax_amount", "line_total"]); workbook["InvoiceLines"].append(["EXT-INVL-001", "EXT-INV-001", "EXT-POL-001", 180, "KG", 100, 3240, 18000])
    workbook["PaymentStatus"].append(["invoice_external_key", "status"]); workbook["PaymentStatus"].append(["EXT-INV-001", "not_released"])
    for name in ("RFQExport", "ComparisonExport", "POExport", "ExceptionExport"): workbook[name].append(["external_key", "status"])
    workbook["SyncLog"].append(["schema_version", "source_version"]); workbook["SyncLog"].append(["1.0", "stale" if stale else "current"])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


if __name__ == "__main__":
    build(ROOT / "genuinegigs_procurement_exchange_v1.xlsx")
    build(ROOT / "conflicts" / "stale_external_version.xlsx", stale=True)
    build(ROOT / "malformed" / "hostile_formula.xlsx", hostile=True)
    invalid = load_workbook(ROOT / "genuinegigs_procurement_exchange_v1.xlsx")
    invalid["Items"].append(["EXT-INVALID-1", "", "Missing material code", "KG"])
    invalid.save(ROOT / "malformed" / "invalid_required_field.xlsx")
    duplicate = load_workbook(ROOT / "genuinegigs_procurement_exchange_v1.xlsx")
    duplicate["Suppliers"].append(["EXT-SUP-001", "Conflicting duplicate supplier", "approved", 10, 10])
    duplicate.save(ROOT / "malformed" / "duplicate_external_key.xlsx")
    partial = Workbook(); partial.active.title = "Items"
    partial["Items"].append(["external_key", "code", "name", "uom"]); partial["Items"].append(["EXT-PARTIAL-1", "PART-1", "", "KG"])
    partial.save(ROOT / "malformed" / "partial_record.xlsx")
    revision = load_workbook(ROOT / "genuinegigs_procurement_exchange_v1.xlsx")
    revision["Items"]["C2"] = "Aluminium Ingot Revised"
    (ROOT / "revisions").mkdir(parents=True, exist_ok=True)
    revision.save(ROOT / "revisions" / "incremental_update_v2.xlsx")
    scale = Workbook(); scale.active.title = "Items"
    scale["Items"].append(["external_key", "code", "name", "uom"])
    for index in range(1, 5001): scale["Items"].append([f"SCALE-{index:05}", f"SC-{index:05}", f"Scale material {index}", "EA"])
    (ROOT / "scale").mkdir(parents=True, exist_ok=True)
    scale.save(ROOT / "scale" / "items_5000_rows.xlsx")
    (ROOT / "mapping_profile_v1.json").write_text(json.dumps({"schema_version": "1.0", "identity": "external_key", "sheets": SHEETS}, indent=2) + "\n")
    (ROOT / "expected_import_summary.json").write_text(json.dumps({"items": 10, "suppliers": 6, "requirements": 1, "invalid_rows": 0}, indent=2) + "\n")
    (ROOT / "expected_export_summary.json").write_text(json.dumps({"write_strategy": "new_version_only", "required_sheets": ["POExport", "ExceptionExport", "SyncLog"]}, indent=2) + "\n")
