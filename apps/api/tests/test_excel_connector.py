from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app import agent_service, excel_connector
from app.agent_schemas import AgentMessageRequest, IntentEnvelope, ThreadCreateRequest, WorkspaceCreateRequest
from app.db import models
from app.db.seed import reset_workspace_database
from app.db.session import SessionLocal
from app.domains import workflows


@pytest.fixture(autouse=True)
def reset_workspace() -> None:
    with SessionLocal() as db:
        reset_workspace_database(db)
        db.commit()


def workbook_bytes(*, item_name: str = "Nickel Ingot", formula: bool = False) -> bytes:
    workbook = Workbook()
    items = workbook.active
    items.title = "Items"
    items.append(["external_key", "code", "name", "uom"])
    items.append(["EXT-NI-001", "NI-001", "=1+1" if formula else item_name, "KG"])
    suppliers = workbook.create_sheet("Suppliers")
    suppliers.append(["external_key", "name", "status", "quality_score", "delivery_score"])
    suppliers.append(["EXT-SUP-001", "Nickel Metals", "approved", 91, 89])
    requirements = workbook.create_sheet("Requirements")
    requirements.append(["external_key", "business_number", "item_code", "quantity", "uom", "need_by_date", "reason"])
    requirements.append(["EXT-REQ-001", "REQ-EXT-0001", "NI-001", 250, "KG", "2026-08-30", "Production raw material"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_fresh_workspace_creator_is_not_duplicated_as_a_demo_seat() -> None:
    with SessionLocal() as db:
        admin = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        tenant, membership = agent_service.create_fresh_workspace(db, admin, WorkspaceCreateRequest(
            company_name="Independent Pilot", workspace_name="Independent Pilot",
            plant_name="Main Plant", plant_code="MAIN", agent_enabled=True,
        ))
        db.flush()
        assert db.query(models.WorkspaceMembership).filter_by(
            tenant_id=tenant.id, account_id=membership.account_id,
        ).count() == 1


def test_xlsx_preview_commit_reimport_and_stale_conflict() -> None:
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user)
        db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(
            tenant_id=user.tenant_id, provider="local",
        ).first()
        content = workbook_bytes()
        preview = excel_connector.preview_import(db, user, connection, "exchange.xlsx", content)
        db.flush()
        assert preview.summary["insert"] == 3
        excel_connector.commit_import(db, user, preview)
        db.flush()
        assert preview.status == "completed"
        assert db.query(models.Item).filter_by(erp_item_code="EXT-NI-001").one().name == "Nickel Ingot"
        assert db.query(models.Supplier).filter_by(erp_vendor_id="EXT-SUP-001").one()
        assert db.query(models.PurchaseRequirement).filter_by(business_number="REQ-EXT-0001").one().source == "excel"

        repeat = excel_connector.preview_import(db, user, connection, "exchange.xlsx", content)
        db.flush()
        assert repeat.summary["unchanged"] == 3
        excel_connector.commit_import(db, user, repeat)
        assert repeat.summary["committed"] == 0

        item = db.query(models.Item).filter_by(erp_item_code="EXT-NI-001").one()
        item.name = "Locally reviewed nickel"
        db.flush()
        stale = excel_connector.preview_import(
            db, user, connection, "exchange.xlsx", workbook_bytes(item_name="External stale nickel"),
        )
        db.flush()
        assert stale.summary["conflict"] == 1
        excel_connector.commit_import(db, user, stale)
        assert item.name == "Locally reviewed nickel"
        assert stale.status == "completed_with_errors"


def test_formula_is_rejected_on_import_and_neutralized_on_export() -> None:
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user)
        db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
        preview = excel_connector.preview_import(db, user, connection, "hostile.xlsx", workbook_bytes(formula=True))
        db.flush()
        assert preview.summary["error"] == 1
        error = db.query(models.IntegrationImportRowResult).filter_by(batch_id=preview.id, action="error").one()
        assert any(message["code"] == "formula_rejected" for message in error.validation_messages)

        supplier = db.query(models.Supplier).first()
        supplier.name = "=HYPERLINK(\"https://invalid.example\")"
        exported = excel_connector.export_exchange_workbook(db, user)
        workbook = load_workbook(BytesIO(exported), data_only=False)
        values = list(workbook["Suppliers"].values)
        exported_name = next(row[1] for row in values[1:] if row[0] == supplier.erp_vendor_id)
        assert exported_name.startswith("'=")


def test_zip_rejects_path_traversal_and_non_csv_members() -> None:
    import zipfile

    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("../escape.csv", "external_key,code,name,uom\n1,A,A,EA")
    with pytest.raises(Exception) as error:
        excel_connector.parse_workbook("bad.zip", output.getvalue())
    assert "safely named" in str(error.value)


def test_golden_workbook_bootstraps_supported_canonical_records_without_manual_sql(monkeypatch: pytest.MonkeyPatch) -> None:
    fixture = Path(__file__).parent / "fixtures" / "excel_erp" / "genuinegigs_procurement_exchange_v1.xlsx"
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user)
        db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
        batch = excel_connector.preview_import(db, user, connection, fixture.name, fixture.read_bytes())
        db.flush()
        excel_connector.commit_import(db, user, batch)
        db.flush()
        assert db.query(models.Company).filter_by(tenant_id=user.tenant_id, erp_code="ORG-APEX").one().name == "Apex Components"
        assert db.query(models.Plant).filter_by(tenant_id=user.tenant_id, erp_location_code="NASHIK").one().name == "Nashik Plant"
        imported_users = db.query(models.User).filter(models.User.tenant_id == user.tenant_id, models.User.email.like("%@fixture.invalid")).all()
        assert {member.role for member in imported_users} == {
            "plant_manager", "purchase_manager", "purchase_executive", "gate_operator",
            "store_manager", "quality_inspector", "admin",
        }
        imported_exec = next(member for member in imported_users if member.role == "purchase_executive")
        imported_manager = next(member for member in imported_users if member.role == "purchase_manager")
        assert imported_exec.manager_id == imported_manager.id
        assert db.query(models.ReportingLine).filter_by(
            tenant_id=user.tenant_id, report_user_id=imported_exec.id, manager_user_id=imported_manager.id,
        ).one()
        imported_account = db.get(models.Account, imported_exec.account_id)
        assert imported_account.status == "invited" and imported_account.email_verified_at is None
        assert db.query(models.Item).filter(models.Item.erp_item_code.like("EXT-ITEM-%")).count() == 10
        assert db.query(models.Supplier).filter(models.Supplier.erp_vendor_id.like("EXT-SUP-%")).count() == 6
        requirement = db.query(models.PurchaseRequirement).filter_by(business_number="REQ-EXT-0001").one()
        capability = db.query(models.SupplierItemCapability).join(models.Item, models.SupplierItemCapability.item_id == models.Item.id).filter(models.Item.code == "MAT-001").one()
        assert capability.approved is True
        assert db.query(models.SupplierContact).filter_by(supplier_id=capability.supplier_id).one().email == "supplier1@fixture.invalid"
        rfq = workflows.create_rfq_from_requirement(db, user, requirement.id); db.flush()
        assert capability.supplier_id in rfq.supplier_ids
        from types import SimpleNamespace
        from app import procurement_v2
        monkeypatch.setattr(
            procurement_v2,
            "generate_artifact",
            lambda *_args, **_kwargs: SimpleNamespace(id="ART-GOLDEN-RFQ", document_id="DOC-GOLDEN-RFQ"),
        )
        publication = workflows.publish_rfq(db, user, rfq.id); db.flush()
        assert publication["status"] == "published"
        message = db.query(models.OutboxMessage).filter_by(
            tenant_id=user.tenant_id, plant_id=user.plant_id, recipient="supplier1@fixture.invalid",
        ).one()
        assert message.status == "approved_pending_send"
        assert message.meta["rfq_id"] == rfq.id
        po = db.query(models.PODraft).filter_by(business_number="PO-EXT-0001").one()
        assert po.award_id is None and po.oracle_mapping["external_owned"] is True
        assert db.query(models.PODraftLine).filter_by(po_draft_id=po.id).one().quantity == 200
        po_detail = workflows.po_draft_payload(db, user, po)
        assert po_detail["provenance"] == "external_system"
        assert po_detail["award"] is None and po_detail["quote"] is None
        receipt = db.query(models.StoreReceipt).filter_by(business_number="GRN-EXT-0001").one()
        assert (receipt.received_quantity, receipt.short_quantity) == (180, 20)
        inspection = db.query(models.InspectionResult).filter_by(business_number="INS-EXT-0001").one()
        assert (inspection.accepted_quantity, inspection.held_quantity) == (170, 10)
        additional_receipt = models.StoreReceipt(
            tenant_id=user.tenant_id, plant_id=user.plant_id, po_draft_id=po.id,
            status="received", received_quantity=1, short_quantity=0, excess_quantity=0, damaged_quantity=0,
        )
        db.add(additional_receipt); db.flush()
        manual_inspection = workflows.record_inspection(db, user, additional_receipt.id, 1, 1, 0, 0); db.flush()
        assert db.query(models.InventoryImpact).filter_by(inspection_id=manual_inspection.id).one().item_id == capability.item_id
        invoice = db.query(models.SupplierInvoice).filter_by(business_number="INV-EXT-0001").one()
        assert db.query(models.SupplierInvoiceLine).filter_by(invoice_id=invoice.id).one().quantity == 180
        payment = db.query(models.InvoicePaymentStatus).filter_by(invoice_id=invoice.id).one()
        assert payment.status == "not_released"
        assert payment.evidence["source"] == "excel_read_only"
        assert batch.status == "completed"
        assert batch.summary["error"] == 0
        repeat = excel_connector.preview_import(db, user, connection, fixture.name, fixture.read_bytes())
        db.flush()
        assert repeat.summary["insert"] == 0
        assert repeat.summary["update"] == 0
        excel_connector.commit_import(db, user, repeat); db.flush()
        assert db.query(models.PODraft).filter_by(business_number="PO-EXT-0001").count() == 1
        assert db.query(models.SupplierInvoice).filter_by(invoice_number="INV-FIX-001").count() == 1

        # Structured inbound exception evidence must survive the same versioned
        # workbook path used for customer-system exchange.
        structured_export = load_workbook(BytesIO(excel_connector.export_exchange_workbook(db, user)))
        receipt_sheet = structured_export["Receipts"]
        receipt_headers = {name: index + 1 for index, name in enumerate(next(receipt_sheet.values))}
        receipt_row = next(
            row for row in range(2, receipt_sheet.max_row + 1)
            if receipt_sheet.cell(row, receipt_headers["external_key"]).value == "EXT-REC-001"
        )
        receipt_sheet.cell(receipt_row, receipt_headers["exception_types"], "damage,missing_certificate")
        receipt_sheet.cell(receipt_row, receipt_headers["observed_item_code"], "MAT-001")
        receipt_sheet.cell(receipt_row, receipt_headers["certificate_status"], "missing")
        receipt_sheet.cell(receipt_row, receipt_headers["exception_notes"], "Crushed packaging on arrival")
        receipt_sheet.cell(receipt_row, receipt_headers["production_impact"], True)
        quality_sheet = structured_export["QualityStatus"]
        quality_headers = {name: index + 1 for index, name in enumerate(next(quality_sheet.values))}
        quality_row = next(
            row for row in range(2, quality_sheet.max_row + 1)
            if quality_sheet.cell(row, quality_headers["external_key"]).value == "EXT-QA-001"
        )
        quality_sheet.cell(quality_row, quality_headers["certificate_status"], "invalid")
        quality_sheet.cell(quality_row, quality_headers["defect_codes"], "SURFACE_DAMAGE,CERT_MISMATCH")
        quality_sheet.cell(quality_row, quality_headers["inspection_notes"], "Certificate heat number does not match")
        quality_sheet.cell(quality_row, quality_headers["production_impact"], True)
        structured_bytes = BytesIO()
        structured_export.save(structured_bytes)
        structured_batch = excel_connector.preview_import(
            db, user, connection, "structured-inbound.xlsx", structured_bytes.getvalue(),
        )
        db.flush(); excel_connector.commit_import(db, user, structured_batch); db.flush()
        db.refresh(receipt); db.refresh(inspection)
        assert receipt.exception_types == ["damage", "missing_certificate"]
        assert receipt.certificate_status == "missing"
        assert receipt.exception_notes == "Crushed packaging on arrival"
        assert receipt.production_impact is True
        assert inspection.certificate_status == "invalid"
        assert inspection.defect_codes == ["SURFACE_DAMAGE", "CERT_MISMATCH"]
        assert inspection.inspection_notes == "Certificate heat number does not match"
        assert inspection.production_impact is True

        po.status = "failed"; db.flush()
        stale_fixture = Path(__file__).parent / "fixtures" / "excel_erp" / "conflicts" / "stale_external_version.xlsx"
        stale = excel_connector.preview_import(db, user, connection, stale_fixture.name, stale_fixture.read_bytes())
        db.flush()
        po_row = db.query(models.IntegrationImportRowResult).filter_by(batch_id=stale.id, sheet_name="OpenPOs").one()
        assert po_row.action == "conflict"
        assert any(message["code"] == "stale_external_data" for message in po_row.validation_messages)

        exported_bytes = excel_connector.export_exchange_workbook(db, user)
        exported = load_workbook(BytesIO(exported_bytes), data_only=False)
        assert {"OpenPOs", "Receipts", "QualityStatus", "Invoices", "InvoiceLines", "PaymentStatus", "Reconciliation", "SyncLog"} <= set(exported.sheetnames)
        assert any(row[0] == "EXT-PO-001" for row in list(exported["OpenPOs"].values)[1:])
        assert any(row[0] == "EXT-INV-001" for row in list(exported["Invoices"].values)[1:])
        assert any(row[0] == "EXT-INV-001" and row[1] == "not_released" for row in list(exported["PaymentStatus"].values)[1:])
        assert any(row[1] == "EXT-PO-001" for row in list(exported["Reconciliation"].values)[1:])
        verification = excel_connector.verify_exchange_export(exported_bytes, db.query(models.IntegrationExternalRecordState).filter_by(connection_id=connection.id).all())
        assert verification["status"] == "reconciled"
        assert verification["verified"] == verification["expected"]
        before = {model.__tablename__: db.query(model).count() for model in (models.Item, models.Supplier, models.PODraft, models.StoreReceipt, models.InspectionResult, models.SupplierInvoice)}
        round_trip = excel_connector.preview_import(db, user, connection, "round-trip.xlsx", exported_bytes)
        db.flush(); excel_connector.commit_import(db, user, round_trip); db.flush()
        after = {model.__tablename__: db.query(model).count() for model in (models.Item, models.Supplier, models.PODraft, models.StoreReceipt, models.InspectionResult, models.SupplierInvoice)}
        assert after == before


def test_versioned_customer_mapping_applies_only_allowlisted_transforms() -> None:
    content = b"Customer Key,Material Number,Material Description,Base Unit\nEXT-CUSTOM-1, ab-77 , Custom Alloy ,kg\n"
    mappings = {"Items": {
        "external_key": "Customer Key", "code": "Material Number",
        "name": "Material Description", "uom": "Base Unit",
    }}
    transforms = {"Items": {"code": "uppercase", "name": "strip", "uom": "uppercase"}}
    excel_connector.validate_mapping_contract(mappings, transforms)
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user); db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
        profile = models.IntegrationMappingProfile(
            tenant_id=user.tenant_id, plant_id=user.plant_id, connection_id=connection.id,
            name="Customer material export", profile_version=1, mappings=mappings,
            transforms=transforms, ownership={"Items": "external_owned"}, status="active",
        )
        db.add(profile); db.flush()
        batch = excel_connector.preview_import(db, user, connection, "materials.csv", content, mapping_profile=profile)
        db.flush()
        excel_connector.commit_import(db, user, batch); db.flush()
        item = db.query(models.Item).filter_by(erp_item_code="EXT-CUSTOM-1").one()
        assert (item.code, item.name, item.uom_id) == ("AB-77", "Custom Alloy", "KG")
        row = db.query(models.IntegrationImportRowResult).filter_by(batch_id=batch.id).one()
        assert row.mapping_version == "1"

    with pytest.raises(Exception, match="Unsupported transforms"):
        excel_connector.validate_mapping_contract(mappings, {"Items": {"name": "python_eval"}})


def test_admin_agent_previews_attached_workbook_through_same_connector(monkeypatch) -> None:
    content = workbook_bytes()
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        document = models.Document(
            tenant_id=user.tenant_id, plant_id=user.plant_id, filename="agent-exchange.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(content), storage_key="test/agent-exchange.xlsx", storage_bucket="evidence",
            checksum_sha256="a" * 64, status="available", validation_errors=[], uploaded_by_user_id=user.id,
        )
        db.add(document); db.flush()
        thread = agent_service.create_thread(db, user, ThreadCreateRequest(thread_type="personal", title="Excel preview"))
        monkeypatch.setattr("app.storage.object_storage.get_bytes", lambda *_args: content)
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (
            IntentEnvelope(intent="preview_excel_import", requested_outcome="Preview this workbook", attachment_ids=[document.id], confidence=1),
            "groq", "test-model", None,
        ))
        run, message = agent_service.run_message(
            db, user, thread.id,
            AgentMessageRequest(content="Check this workbook before importing it", attachment_ids=[document.id]),
        )
        assert run.intent == "preview_excel_import"
        assert "Workbook preview ready" in message.content
        assert db.query(models.IntegrationImportBatch).filter_by(document_id=document.id, status="previewed").one()
        assert db.query(models.AgentActionReceipt).filter_by(tool_name="preview_excel_import", status="succeeded").one()


@pytest.mark.parametrize(("filename", "error_code"), [
    ("invalid_required_field.xlsx", "required"),
    ("partial_record.xlsx", "required"),
    ("duplicate_external_key.xlsx", "duplicate_external_key"),
    ("hostile_formula.xlsx", "formula_rejected"),
])
def test_negative_fixture_rows_are_isolated_with_specific_errors(filename: str, error_code: str) -> None:
    fixture = Path(__file__).parent / "fixtures" / "excel_erp" / "malformed" / filename
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user); db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
        batch = excel_connector.preview_import(db, user, connection, filename, fixture.read_bytes()); db.flush()
        codes = {message["code"] for row in db.query(models.IntegrationImportRowResult).filter_by(batch_id=batch.id, action="error").all() for message in row.validation_messages}
        assert error_code in codes


def test_scale_fixture_discovers_all_rows_with_bounded_parser() -> None:
    fixture = Path(__file__).parent / "fixtures" / "excel_erp" / "scale" / "items_5000_rows.xlsx"
    discovery = excel_connector.discover(fixture.name, fixture.read_bytes())
    assert discovery["sheets"][0]["row_count"] == 5000


def test_zip_csv_parser_rejects_archive_fanout_and_path_traversal() -> None:
    too_many = BytesIO()
    with zipfile.ZipFile(too_many, "w", zipfile.ZIP_DEFLATED) as archive:
        for index in range(excel_connector.MAX_ARCHIVE_FILES + 1):
            archive.writestr(
                f"sheet-{index}.csv",
                "external_key,code,name,uom\n"
                f"EXT-{index},I-{index},Item {index},EA\n",
            )
    with pytest.raises(Exception) as fanout:
        excel_connector.parse_workbook("oversized.zip", too_many.getvalue())
    assert getattr(fanout.value, "status_code", None) == 422
    assert "safe extraction limits" in str(fanout.value)

    traversal = BytesIO()
    with zipfile.ZipFile(traversal, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "../items.csv",
            "external_key,code,name,uom\nEXT-1,I-1,Unsafe path,EA\n",
        )
    with pytest.raises(Exception) as unsafe_path:
        excel_connector.parse_workbook("traversal.zip", traversal.getvalue())
    assert getattr(unsafe_path.value, "status_code", None) == 422
    assert "safely named CSV" in str(unsafe_path.value)


def test_import_reconciliation_report_is_row_level_and_formula_safe() -> None:
    content = b"external_key,code,name,uom\nEXT-REPORT-1,R-1,=HYPERLINK(\"\"https://bad.invalid\"\"),EA\n"
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user); db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
        batch = excel_connector.preview_import(db, user, connection, "report.csv", content); db.flush()
        report = excel_connector.import_reconciliation_report(db, user, batch).decode("utf-8-sig")
        assert "source_file,source_hash,batch_status,sheet,row,external_key,action" in report
        assert "formula_rejected" in report
        assert "=HYPERLINK" not in report


def test_admin_agent_export_uses_same_versioned_service_and_persists_receipt(monkeypatch: pytest.MonkeyPatch) -> None:
    stored: list[bytes] = []
    monkeypatch.setattr("app.storage.object_storage.put_file", lambda _bucket, _key, body, _type: stored.append(body.read()))
    monkeypatch.setattr("app.storage.object_storage.presigned_get", lambda _bucket, key, _ttl: f"/signed/{key}")
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        thread = agent_service.create_thread(db, user, ThreadCreateRequest(thread_type="personal", title="Excel export"))
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (
            IntentEnvelope(intent="export_excel_exchange", requested_outcome="Export the current procurement exchange", confidence=1),
            "groq", "test-model", None,
        ))
        run, message = agent_service.run_message(db, user, thread.id, AgentMessageRequest(content="Create a fresh Excel export"))
        batch = db.query(models.IntegrationExportBatch).order_by(models.IntegrationExportBatch.created_at.desc()).first()
        assert run.intent == "export_excel_exchange"
        assert batch is not None and batch.status == "completed" and batch.output_version >= 1
        assert stored and stored[0].startswith(b"PK")
        assert batch.summary["read_after_write"]["status"] == "reconciled"
        assert db.query(models.AgentActionReceipt).filter_by(
            run_id=run.id, tool_name="export_excel_exchange", target_entity_id=batch.id,
        ).one()
        assert batch.summary["filename"] in message.content


def test_admin_agent_import_requires_confirmation_before_canonical_commit(monkeypatch: pytest.MonkeyPatch) -> None:
    content = b"external_key,code,name,uom\nEXT-CONFIRM-1,CONF-1,Confirmation Alloy,KG\n"
    with SessionLocal() as db:
        user = db.query(models.User).filter_by(email="admin@genuinegigs.local").one()
        workflows.ensure_default_connections(db, user); db.flush()
        connection = db.query(models.IntegrationConnection).filter_by(provider="local").first()
        batch = excel_connector.preview_import(db, user, connection, "confirmation.csv", content); db.flush()
        thread = agent_service.create_thread(db, user, ThreadCreateRequest(thread_type="personal", title="Excel approval"))
        monkeypatch.setattr(agent_service, "provider_response", lambda *_args, **_kwargs: (
            IntentEnvelope(
                intent="approve_excel_import_proposal", requested_outcome="Approve the clean workbook",
                entities={"batch_id": batch.id}, confidence=1,
            ),
            "groq", "test-model", None,
        ))
        run, _message = agent_service.run_message(db, user, thread.id, AgentMessageRequest(content="Approve that clean import"))
        proposal = db.query(models.AgentProposal).filter_by(run_id=run.id, action="approve_excel_import").one()
        assert proposal.confirmation_state == "pending"
        assert db.query(models.Item).filter_by(erp_item_code="EXT-CONFIRM-1").count() == 0
        agent_service.confirm_proposal(db, user, proposal.id); db.flush()
        assert db.query(models.Item).filter_by(erp_item_code="EXT-CONFIRM-1").one().code == "CONF-1"
        assert batch.status == "completed"
